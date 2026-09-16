"""Serviço de importação de entidades e contratos a partir de ficheiros ZIP, XLSX e JSON."""
import json
import re
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from api.elasticsearch_client import (
    delete_contracts_by_ids,
    ensure_indices,
    find_contract_ids_by_idcontrato,
    get_es_client,
    index_contracts,
)
from collectors.contratos import normalize_contract

ROOT = Path(__file__).resolve().parents[1]
ENTITIES_PATH = ROOT / "data" / "entidades-gov-portal-base" / "entidades.json"

KNOWN_CONTRACT_FIELDS = {
    "idcontrato", "idprocedimento", "objectoContrato", "descContrato",
    "adjudicante", "adjudicantes", "adjudicatario", "adjudicatarios",
    "precoContratual", "dataCelebracaoContrato", "dataPublicacao",
    "Ano", "tipoContrato", "tipoprocedimento", "cpv", "NUTs",
    "localExecucao", "prazoExecucao",
}
KNOWN_ENTITY_FIELDS = {
    "nif", "nifEntidade", "name", "desigEntidade", "nome", "denominacao",
    "numContratos", "totAdjudicante", "totAdjudicatario", "totValorContratIni",
    "pais", "descPais", "AliasPais",
}


def _is_valid_nif(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    return bool(re.fullmatch(r"\d{9}", text))


def _first_value(row: Dict[str, Any], keys: List[str]) -> Any:
    for k in keys:
        if k in row and row[k] not in (None, "", "-"):
            return row[k]
    return None


def _parse_date(value: Any) -> Optional[str]:
    if not value or not isinstance(value, str):
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value.strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def _load_entities_lookup() -> Dict[str, Dict[str, Any]]:
    """Carrega entidades conhecidas do portal base (ficheiro local) indexadas por NIF."""
    lookup: Dict[str, Dict[str, Any]] = {}
    if not ENTITIES_PATH.exists():
        return lookup
    try:
        data = json.loads(ENTITIES_PATH.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            data = [data]
        for item in data or []:
            nif = str(item.get("nifEntidade") or item.get("nif") or "").strip()
            if not _is_valid_nif(nif):
                continue
            lookup[nif] = {
                "nif": nif,
                "name": item.get("desigEntidade") or item.get("name") or item.get("nome") or "",
                "country": item.get("descPais") or item.get("AliasPais") or item.get("pais") or "",
                "contracts_count": item.get("numContratos") or 0,
                "total_value": item.get("totValorContratIni") or 0,
                "adjudicante_count": item.get("totAdjudicante") or 0,
                "adjudicatario_count": item.get("totAdjudicatario") or 0,
                "source": "local_entities_file",
            }
    except Exception:
        pass
    return lookup


def _detect_contract_vs_entity(rows: List[Dict[str, Any]]) -> str:
    """Heurística simples para distinguir entidades de contratos."""
    if not rows:
        return "entities"
    keys = set(rows[0].keys())
    contract_score = len(keys & KNOWN_CONTRACT_FIELDS)
    entity_score = len(keys & KNOWN_ENTITY_FIELDS)
    if contract_score > entity_score:
        return "contracts"
    if entity_score > contract_score:
        return "entities"
    # fallback por conteúdo
    for k in ["objectoContrato", "precoContratual", "dataCelebracaoContrato"]:
        if k in keys:
            return "contracts"
    return "entities"


def _load_json_data(content: bytes) -> List[Dict[str, Any]]:
    data = json.loads(content.decode("utf-8-sig"))
    if isinstance(data, dict):
        for k in ["contratos", "data", "items", "entities", "entidades", "results"]:
            if k in data and isinstance(data[k], list):
                return data[k]
        return [data]
    if isinstance(data, list):
        return data
    return []


def _rows_from_xlsx(content: bytes) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl não disponível") from exc

    wb = load_workbook(filename=BytesIO(content), data_only=True)
    rows: List[Dict[str, Any]] = []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(values[0])]
        for r in values[1:]:
            row = {headers[i]: v for i, v in enumerate(r) if i < len(headers)}
            if any(v is not None and str(v).strip() != "" for v in row.values()):
                rows.append(row)
    return rows


def _extract_json_from_zip(content: bytes) -> List[Dict[str, Any]]:
    with zipfile.ZipFile(BytesIO(content), "r") as zf:
        names = [n for n in zf.namelist() if n.lower().endswith((".json", ".jsonl", ".xlsx"))]
        if not names:
            raise ValueError("ZIP não contém .json, .jsonl ou .xlsx")
        rows: List[Dict[str, Any]] = []
        for name in names:
            data = zf.read(name)
            if name.lower().endswith(".xlsx"):
                rows.extend(_rows_from_xlsx(data))
            else:
                rows.extend(_load_json_data(data))
        return rows


def _extract_text_and_nif(value: Any) -> Dict[str, Any]:
    text = ""
    nif = ""
    if value is None:
        return {"text": text, "nif": nif, "parsed": []}
    if isinstance(value, str):
        text = value
    elif isinstance(value, list):
        text = "\n".join(str(v).strip() for v in value if v is not None)
    else:
        text = str(value)
    parsed = []
    for m in re.finditer(r"(\d{9})\s*-\s*([^,\n]+)", text):
        parsed.append({"nif": m.group(1), "nome": m.group(2).strip()})
    # se não houver padrão NIF - NOME, procurar NIF isolado e manter texto
    if not parsed:
        for m in re.finditer(r"(?<![0-9])\d{9}(?![0-9])", text):
            parsed.append({"nif": m.group(0), "nome": text.strip()[:120]})
    return {"text": text, "nif": nif, "parsed": parsed}


def _build_preview_row_contract(raw: Dict[str, Any]) -> Dict[str, Any]:
    id_ = str(_first_value(raw, ["idcontrato", "idprocedimento", "id"]) or "")
    obj = str(_first_value(raw, ["objectoContrato", "objecto", "objeto", "description"]) or "")
    preco = _as_float(_first_value(raw, ["precoContratual", "preco", "price", "valor"]))
    data = _parse_date(_first_value(raw, ["dataCelebracaoContrato", "dataCelebracao", "data", "date"]))

    adj_raw = _extract_text_and_nif(_first_value(raw, ["adjudicante", "adjudicantes"]))
    adj_parsed = adj_raw["parsed"]
    adj_text = "; ".join(f"{p['nif']} - {p['nome']}" for p in adj_parsed) or adj_raw["text"]

    adjd_raw = _extract_text_and_nif(_first_value(raw, ["adjudicatario", "adjudicatarios", "adjudicatários"]))
    adjd_parsed = adjd_raw["parsed"]
    adjd_text = "; ".join(f"{p['nif']} - {p['nome']}" for p in adjd_parsed) or adjd_raw["text"]

    return {
        "id": id_,
        "name": None,
        "nif": None,
        "objectoContrato": obj,
        "precoContratual": preco,
        "dataCelebracaoContrato": data,
        "adjudicante": adj_text,
        "adjudicatario": adjd_text,
        "raw": raw,
    }


def _build_preview_row_entity(raw: Dict[str, Any]) -> Dict[str, Any]:
    nif = str(_first_value(raw, ["nifEntidade", "nif", "NIF"]) or "")
    name = str(_first_value(raw, ["desigEntidade", "name", "nome", "denominacao", "designacao"]) or "")
    return {
        "id": nif,
        "name": name,
        "nif": nif,
        "objectoContrato": None,
        "precoContratual": _as_float(raw.get("totValorContratIni")),
        "dataCelebracaoContrato": None,
        "adjudicante": None,
        "adjudicatario": None,
        "raw": raw,
    }


def parse_file(filename: str, content: bytes, data_type: str = "auto") -> Dict[str, Any]:
    """Analisa um ficheiro e devolve metadados + rows de pré-visualização."""
    lower = filename.lower()
    if lower.endswith(".zip"):
        file_type = "zip"
        raw_rows = _extract_json_from_zip(content)
    elif lower.endswith(".xlsx"):
        file_type = "xlsx"
        raw_rows = _rows_from_xlsx(content)
    elif lower.endswith(".json") or lower.endswith(".jsonl"):
        file_type = "json"
        raw_rows = _load_json_data(content)
    else:
        raise ValueError(f"Formato não suportado: {filename}")

    detected = _detect_contract_vs_entity(raw_rows) if data_type == "auto" else data_type
    if detected not in ("contracts", "entities"):
        detected = "contracts"

    warnings: List[str] = []
    if data_type != "auto" and data_type != detected:
        warnings.append(f"Tipo pedido ({data_type}) difere do detetado ({detected}); a usar {data_type}.")
        detected = data_type

    rows: List[Dict[str, Any]] = []
    for raw in raw_rows:
        try:
            if detected == "contracts":
                rows.append(_build_preview_row_contract(raw))
            else:
                rows.append(_build_preview_row_entity(raw))
        except Exception as e:
            warnings.append(f"Linha ignorada: {e}")

    schema: List[str] = []
    if raw_rows:
        schema = sorted(raw_rows[0].keys())[:40]

    return {
        "data_type": detected,
        "file_type": file_type,
        "filename": filename,
        "rows": rows,
        "total_rows": len(rows),
        "sample_schema": schema,
        "errors": [],
        "warnings": warnings,
    }


def _normalize_preview_contract(row: Dict[str, Any], link_entities: bool = True) -> Dict[str, Any]:
    raw = row.get("raw", {})
    # Constrói um registo bruto compatível com normalize_contract
    nif_adjudicante = row.get("nif") if _is_valid_nif(row.get("nif")) else None
    adjudicante_text = row.get("adjudicante") or ""
    adjudicatario_text = row.get("adjudicatario") or ""

    # Se tivermos um NIF isolado e adjudicante vazio, preenche adjudicante
    if nif_adjudicante and not adjudicante_text and row.get("name"):
        adjudicante_text = f"{nif_adjudicante} - {row.get('name')}"

    raw_contract = {
        "idcontrato": row.get("id") or raw.get("idcontrato") or raw.get("idprocedimento") or f"import-{datetime.utcnow().timestamp()}",
        "objectoContrato": row.get("objectoContrato") or raw.get("objectoContrato") or raw.get("objecto") or "",
        "descContrato": raw.get("descContrato") or "",
        "adjudicante": adjudicante_text or raw.get("adjudicante") or raw.get("adjudicantes") or "",
        "adjudicatarios": adjudicatario_text or raw.get("adjudicatario") or raw.get("adjudicatarios") or raw.get("adjudicatários") or "",
        "precoContratual": row.get("precoContratual") if row.get("precoContratual") is not None else raw.get("precoContratual") or raw.get("preco") or raw.get("valor"),
        "dataCelebracaoContrato": row.get("dataCelebracaoContrato") or raw.get("dataCelebracaoContrato") or raw.get("dataCelebracao") or raw.get("data") or raw.get("date"),
        "dataPublicacao": raw.get("dataPublicacao") or raw.get("dataPublicacao") or row.get("dataCelebracaoContrato"),
        "Ano": raw.get("Ano") or (row.get("dataCelebracaoContrato")[:4] if row.get("dataCelebracaoContrato") else None),
        "tipoContrato": raw.get("tipoContrato") or raw.get("tipo") or "",
        "tipoprocedimento": raw.get("tipoprocedimento") or raw.get("procedimento") or "",
        "cpv": raw.get("cpv") or "",
        "NUTs": raw.get("NUTs") or raw.get("nuts") or raw.get("regiao") or "",
        "localExecucao": raw.get("localExecucao") or raw.get("local") or "",
        "prazoExecucao": raw.get("prazoExecucao") or raw.get("prazo") or "",
    }
    doc = normalize_contract(raw_contract)

    if link_entities:
        entities_lookup = _load_entities_lookup()
        for party_list in [doc.get("adjudicantes", {}).get("parsed", []), doc.get("adjudicatarios", {}).get("parsed", [])]:
            for party in party_list:
                nif = party.get("nif")
                if nif and nif in entities_lookup:
                    party["nome"] = party.get("nome") or entities_lookup[nif].get("name", "")
        # Atualiza entidades extraídas
        entities = []
        for a in doc["adjudicantes"]["parsed"]:
            entities.append({"name": a["nome"], "type": "adjudicante", "nif": a["nif"]})
        for a in doc["adjudicatarios"]["parsed"]:
            entities.append({"name": a["nome"], "type": "adjudicatario", "nif": a["nif"]})
        for c in doc.get("cpv", []):
            if c.get("description"):
                entities.append({"name": c["description"], "type": "cpv", "code": c["code"]})
        doc["entities"] = entities

    return doc


def _normalize_preview_entity(row: Dict[str, Any]) -> Dict[str, Any]:
    nif = str(row.get("nif") or "").strip()
    name = str(row.get("name") or "").strip()
    raw = row.get("raw", {})
    return {
        "nif": nif,
        "name": name or raw.get("desigEntidade") or raw.get("name") or raw.get("nome") or "",
        "country": raw.get("descPais") or raw.get("AliasPais") or raw.get("pais") or "",
        "contracts_count": raw.get("numContratos") or 0,
        "adjudicante_count": raw.get("totAdjudicante") or 0,
        "adjudicatario_count": raw.get("totAdjudicatario") or 0,
        "total_value": _as_float(raw.get("totValorContratIni")),
        "source": "imported_entities",
        "imported_at": datetime.utcnow().isoformat(),
    }


def ingest_rows(
    rows: List[Dict[str, Any]],
    data_type: str,
    options: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Indexa as linhas de pré-visualização (contratos no ES; entidades em memória/disco).

    Deteção de duplicados:
    - Se algum idcontrato já existir no Elasticsearch e hard_reprocess=False, devolve aviso sem indexar.
    - Se hard_reprocess=True, apaga os documentos existentes e reindexa.
    """
    opts = options or {}
    link_entities = bool(opts.get("link_entities", True))
    max_records = opts.get("max_records")
    skip_validation = bool(opts.get("skip_validation", False))
    hard_reprocess = bool(opts.get("hard_reprocess", False))

    es = get_es_client()
    if not es:
        return {"success": False, "error": "Elasticsearch indisponível", "indexed_count": 0, "total": 0}

    ensure_indices(es)

    total = 0
    indexed = 0
    errors = 0
    linked_entities = 0
    docs: List[Dict[str, Any]] = []
    imported_entities: List[Dict[str, Any]] = []

    for row in rows:
        if max_records and total >= max_records:
            break
        try:
            if data_type == "contracts":
                doc = _normalize_preview_contract(row, link_entities=link_entities)
                docs.append(doc)
                linked_entities += len(doc.get("entities", []))
            else:
                entity = _normalize_preview_entity(row)
                imported_entities.append(entity)
        except Exception as e:
            errors += 1
            if not skip_validation:
                return {"success": False, "error": f"Erro ao normalizar linha {total}: {e}", "indexed_count": 0, "total": total}
        total += 1

    if data_type == "contracts":
        if not docs:
            return {
                "success": True,
                "indexed_count": 0,
                "total": total,
                "errors": errors,
                "linked_entities": 0,
                "duplicate_count": 0,
                "deleted_count": 0,
                "duplicate_ids": [],
                "message": "Nenhum contrato para indexar",
            }

        idcontratos = [str(doc.get("idcontrato") or doc.get("idprocedimento")) for doc in docs if doc.get("idcontrato") or doc.get("idprocedimento")]
        existing_ids = find_contract_ids_by_idcontrato(idcontratos, es) if idcontratos else []

        if existing_ids:
            if not hard_reprocess:
                sample = existing_ids[:50]
                return {
                    "success": False,
                    "indexed_count": 0,
                    "total": total,
                    "errors": errors,
                    "linked_entities": 0,
                    "duplicate_count": len(existing_ids),
                    "deleted_count": 0,
                    "duplicate_ids": sample,
                    "message": f"{len(existing_ids)} contrato(s) já processado(s). Ativa \"Hard reprocess\" para apagar e reindexar.",
                }
            del_res = delete_contracts_by_ids(existing_ids, es)
            if del_res.get("error"):
                return {
                    "success": False,
                    "error": del_res["error"],
                    "indexed_count": 0,
                    "total": total,
                    "errors": errors,
                    "linked_entities": 0,
                    "duplicate_count": len(existing_ids),
                    "deleted_count": del_res.get("deleted_count", 0),
                    "duplicate_ids": existing_ids[:50],
                }

        res = index_contracts(docs, es)
        indexed = res.get("indexed_count", 0)
        errors += res.get("errors", 0)
        if res.get("error"):
            return {
                "success": False,
                "error": res["error"],
                "indexed_count": indexed,
                "total": total,
                "errors": errors,
                "linked_entities": linked_entities,
                "duplicate_count": len(existing_ids),
                "deleted_count": del_res.get("deleted_count", 0) if existing_ids and hard_reprocess else 0,
                "duplicate_ids": existing_ids[:50],
            }
        return {
            "success": True,
            "indexed_count": indexed,
            "total": total,
            "errors": errors,
            "linked_entities": linked_entities,
            "duplicate_count": len(existing_ids),
            "deleted_count": del_res.get("deleted_count", 0) if existing_ids and hard_reprocess else 0,
            "duplicate_ids": existing_ids[:50],
            "message": f"{indexed} contratos indexados de {total}" + (f" ({len(existing_ids)} duplicados apagados)" if existing_ids else ""),
        }

    # Entidades: guardar num ficheiro JSONL de entidades importadas
    dest = ROOT / "data" / "processed" / "entities" / "imported_entities.jsonl"
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "a", encoding="utf-8") as fh:
        for e in imported_entities:
            fh.write(json.dumps(e, ensure_ascii=False) + "\n")

    return {
        "success": True,
        "indexed_count": len(imported_entities),
        "total": len(imported_entities),
        "errors": errors,
        "linked_entities": 0,
        "message": f"{len(imported_entities)} entidades importadas",
        "details": {"destination": str(dest)},
    }
