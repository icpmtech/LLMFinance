"""Visualizador — motor de BI do IQ OS (dados → dimensões, medidas e visuais).

Este módulo transforma tudo o que a plataforma já tem indexado num **modelo
semântico** ao estilo Power BI:

* **datasets** — cada tipo de objeto da Ontologia (`contrato`, `noticia`,
  `marca`, `conta`, …) é um dataset consultável, com as suas dimensões,
  medidas e filtros declarados;
* **medidas** — contagem, soma, média, mínimo, máximo, distintos, mediana e
  percentil 90, calculadas por agregações do Elasticsearch (nunca em memória
  sobre amostras, exceto nos datasets locais — ver abaixo);
* **fórmulas** — medidas calculadas definidas pelo utilizador
  (`[Valor contratual] / [Contagem]`), avaliadas sobre cada linha do resultado;
* **datasets locais** — fontes que não vivem no Elasticsearch (Office, dossiês
  360, contas de Email) são carregadas para memória e agregadas em Python, com
  as mesmas medidas e fórmulas.

O motor é deliberadamente declarativo: a definição dos datasets vem da
Ontologia (`api/ontology_registry.py`), pelo que adicionar um tipo novo à
ontologia torna-o automaticamente analisável aqui — sem código.

Notas de honestidade dos dados (as mesmas que o resto da plataforma):

* agregações de **uma** dimensão usam `terms`/`date_histogram` com `size`
  limitado: quando o número de buckets bate no limite, a resposta é marcada
  como truncada (`meta.truncated`) em vez de esconder o facto;
* `terms` com `missing` cria um grupo «Não especificado» — os contratos sem
  NUTS aparecem, não desaparecem;
* o dataset `contrato` tem campos **aninhados** (`adjudicantes`, `cpv`): as
  transições entre níveis aninhados são feitas com `nested`/`reverse_nested`,
  para que a contagem continue a ser de contratos e não de partes.
"""
from __future__ import annotations

import ast
import csv
import io
import logging
import math
import re
import threading
import time
import unicodedata
from datetime import datetime, timezone
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from api import ontology_registry as registry
from api import ontology_service as ontology
from api.elasticsearch_client import get_es_client

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Limites (documentados na resposta para a UI poder avisar)
# ---------------------------------------------------------------------------
MAX_DIMENSIONS = 3
MAX_MEASURES = 12
MAX_FORMULAS = 12
MAX_ROWS = 1000
TERMS_MAX = 1000
DEFAULT_LIMIT = 25
AGG_TIMEOUT_SECONDS = 45
MAX_LOCAL_ROWS = 20000

_mapping_cache: Dict[str, Tuple[float, Dict[str, Any]]] = {}
_mapping_lock = threading.Lock()
_MAPPING_TTL = 600

#: Tipos de medida suportados e se são somáveis entre linhas (usado no grupo
#: «Outros» do Top-N, onde somar uma mediana ou um cardinalidade seria mentira).
MEASURE_KINDS: Dict[str, Dict[str, Any]] = {
    "contagem": {"label": "Contagem", "additive": True, "field": False},
    "soma": {"label": "Soma", "additive": True, "field": True},
    "media": {"label": "Média", "additive": False, "field": True},
    "minimo": {"label": "Mínimo", "additive": False, "field": True},
    "maximo": {"label": "Máximo", "additive": False, "field": True},
    "distintos": {"label": "Valores distintos", "additive": False, "field": True},
    "mediana": {"label": "Mediana", "additive": False, "field": True},
    "p90": {"label": "Percentil 90", "additive": False, "field": True},
}

INTERVALS: Dict[str, str] = {
    "dia": "day",
    "semana": "week",
    "mes": "month",
    "trimestre": "quarter",
    "ano": "year",
}

NUMERIC_ES_TYPES = {
    "integer", "long", "short", "byte", "double", "float", "half_float",
    "scaled_float", "unsigned_long",
}
DIRECT_AGG_TYPES = NUMERIC_ES_TYPES | {"keyword", "date", "boolean", "ip", "flattened", "date_nanos"}

FORMAT_BY_UNIT = {"€": "currency", "contratos": "integer", "registos": "integer", "notícias": "integer"}

#: Medidas com nome de negócio (mais legíveis do que `soma_preco`) para os
#: datasets agregados, onde a Ontologia já declara o que faz sentido medir.
CURATED_MEASURES: Dict[str, List[Dict[str, str]]] = {
    "cpv": [
        {"id": "valor", "label": "Valor contratual (€)", "kind": "soma", "prop": "preco"},
        {"id": "valor_medio", "label": "Valor médio (€)", "kind": "media", "prop": "preco"},
        {"id": "valor_maximo", "label": "Valor máximo (€)", "kind": "maximo", "prop": "preco"},
    ],
    "regiao": [
        {"id": "valor", "label": "Valor contratual (€)", "kind": "soma", "prop": "preco"},
        {"id": "valor_medio", "label": "Valor médio (€)", "kind": "media", "prop": "preco"},
        {"id": "valor_maximo", "label": "Valor máximo (€)", "kind": "maximo", "prop": "preco"},
    ],
    "ticker": [
        {"id": "cotacoes", "label": "Cotações", "kind": "contagem", "prop": None},
        {"id": "volume_total", "label": "Volume total", "kind": "soma", "prop": "volume"},
        {"id": "fecho_medio", "label": "Fecho médio (€)", "kind": "media", "prop": "fecho"},
        {"id": "fecho_maximo", "label": "Fecho máximo (€)", "kind": "maximo", "prop": "fecho"},
        {"id": "fecho_minimo", "label": "Fecho mínimo (€)", "kind": "minimo", "prop": "fecho"},
        {"id": "primeira_data", "label": "Primeira cotação", "kind": "minimo", "prop": "data"},
        {"id": "ultima_data", "label": "Última cotação", "kind": "maximo", "prop": "data"},
    ],
    "topico": [
        {"id": "distintos_titulos", "label": "Títulos distintos", "kind": "distintos", "prop": "titulo"},
    ],
}

#: Dimensões com nome de negócio para datasets agregados (a chave é a dimensão
#: primária do próprio dataset).
CURATED_DIMENSIONS: Dict[str, Dict[str, str]] = {
    "cpv": {"id": "cpv", "label": "Código CPV", "prop": "codigo"},
    "regiao": {"id": "regiao", "label": "Região (NUTS)", "prop": "nome"},
    "ticker": {"id": "ticker", "label": "Instrumento", "prop": "simbolo"},
    "topico": {"id": "topico", "label": "Tópico", "prop": "nome"},
}


# ---------------------------------------------------------------------------
# Utilitários
# ---------------------------------------------------------------------------
def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize(value: Optional[str]) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).strip().lower()


def _to_number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _as_list(value: Any) -> List[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _prop_map(obj_type: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {prop["id"]: prop for prop in obj_type.get("properties") or []}


def _pk_id(obj_type: Dict[str, Any]) -> Optional[str]:
    for prop in obj_type.get("properties") or []:
        if prop.get("pk"):
            return prop["id"]
    return obj_type.get("primary_key")


# ---------------------------------------------------------------------------
# Mapeamento do índice (para saber como agregar cada campo)
# ---------------------------------------------------------------------------
def _index_properties(client: Any, index: str) -> Dict[str, Any]:
    """Propriedades do mapping do índice (com cache curta)."""
    now = time.time()
    with _mapping_lock:
        cached = _mapping_cache.get(index)
        if cached and now - cached[0] < _MAPPING_TTL:
            return cached[1]
    if not client or not index:
        return {}
    try:
        mapping = client.indices.get_mapping(index=index)
        first = next(iter(mapping.values()), {})
        props = first.get("mappings", {}).get("properties", {}) or {}
    except Exception as exc:  # índice inexistente, ES em baixo, etc.
        logger.debug("Mapping indisponível para %s: %s", index, exc)
        props = {}
    with _mapping_lock:
        _mapping_cache[index] = (now, props)
    return props


def _mapping_spec(properties: Dict[str, Any], field: str) -> Dict[str, Any]:
    """Especificação do campo, atravessando objetos aninhados (`a.b.c`)."""
    node: Any = properties
    for part in str(field or "").split("."):
        if not isinstance(node, dict):
            return {}
        spec = node.get(part)
        if spec is None:
            return {}
        if "properties" in spec:
            node = spec["properties"]
            continue
        return spec
    return {}


def _script_emit(field: str) -> str:
    """Script de runtime field que emite um campo de `_source` (aceita listas)."""
    source = (
        "def v = params._source;"
        + "".join(
            f" if (v != null && v.containsKey('{part}')) {{ v = v.get('{part}'); }} else {{ v = null; }}"
            for part in str(field).split(".")
        )
        + " if (v == null) { return; }"
        " if (v instanceof List) { for (def item : v) { if (item != null) { emit(item); } } } else { emit(v); }"
    )
    return source


def _agg_target(client: Any, index: str, field: str, properties: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Como agregar um campo: direto, `.keyword` ou runtime field.

    O mapping manda: campos declarados como `text` sem subcampo `.keyword` não
    são agregáveis, pelo que se cria um runtime field (mesma estratégia já usada
    pelo construtor de grafos dos contratos).
    """
    props = properties if properties is not None else _index_properties(client, index)
    spec = _mapping_spec(props, field)
    if spec.get("type") in DIRECT_AGG_TYPES:
        return {"field": field, "runtime": None, "type": spec.get("type")}
    if (spec.get("fields") or {}).get("keyword"):
        return {"field": f"{field}.keyword", "runtime": None, "type": "keyword"}
    if not spec and "." in str(field):
        # Campo de objeto aninhado desconhecido: assumir keyword do mapping real.
        return {"field": field, "runtime": None, "type": "keyword"}
    if not spec:
        return {"field": field, "runtime": None, "type": "unknown"}
    if spec.get("type") == "text" and "." in str(field):
        # Runtime fields não veem o interior de `nested`: agrega-se o campo tal
        # como está (o mapping costuma ter keyword nessas folhas).
        return {"field": field, "runtime": None, "type": "keyword", "note": "campo aninhado sem keyword declarado"}
    name = re.sub(r"[^a-zA-Z0-9_]", "_", f"{field}_kw")
    return {
        "field": name,
        "runtime": {name: {"type": "keyword", "script": {"source": _script_emit(field)}}},
        "type": "keyword",
        "note": "agregado por runtime field (o campo é texto sem subcampo keyword)",
    }


# ---------------------------------------------------------------------------
# Catálogo de datasets
# ---------------------------------------------------------------------------
def _field_entry(prop: Dict[str, Any], obj_type: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Dimensão/filtro a partir de uma propriedade da Ontologia."""
    if not prop.get("field"):
        return None
    if prop.get("type") not in ("keyword", "enum", "number", "date", "boolean", "text"):
        return None
    return {
        "id": prop["id"],
        "label": prop.get("label") or prop["id"],
        "type": prop.get("type"),
        "field": prop["field"],
        "nested": prop.get("nested"),
        "unit": prop.get("unit"),
        "enum": prop.get("enum"),
        "values_from": prop.get("values_from"),
        "sortable": bool(prop.get("sortable")),
        "searchable": bool(prop.get("searchable")),
        "filterable": bool(prop.get("filterable", True)),
        "pk": bool(prop.get("pk")),
        "groupable": prop.get("type") != "boolean",
        "from": obj_type["id"],
    }


def _measures_for(obj_type: Dict[str, Any], props: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Medidas derivadas das propriedades numéricas e de data do tipo."""
    out: List[Dict[str, Any]] = []
    for prop in obj_type.get("properties") or []:
        if not prop.get("field") or prop.get("pk"):
            continue
        ptype = prop.get("type")
        unit = prop.get("unit")
        base = {"field": prop["field"], "nested": prop.get("nested"), "unit": unit,
                "format": FORMAT_BY_UNIT.get(str(unit), "number" if ptype == "number" else None),
                "prop": prop["id"], "source": obj_type["id"]}
        if ptype == "number":
            # Anos não se somam: só fazem sentido como dimensão (`Soma de Ano`
            # seria uma medida sem significado).
            if prop["id"] in ("ano", "primeiro_ano", "ultimo_ano") or prop["id"].endswith("_ano"):
                continue
            for kind, label in (("soma", "Soma de"), ("media", "Média de"), ("minimo", "Mínimo de"), ("maximo", "Máximo de")):
                out.append({**base, "id": f"{kind}_{prop['id']}", "label": f"{label} {prop.get('label') or prop['id']}", "kind": kind})
        elif ptype == "date":
            for kind, label in (("minimo", "Primeira"), ("maximo", "Última")):
                out.append({**base, "id": f"{kind}_{prop['id']}", "label": f"{label} {prop.get('label') or prop['id']}", "kind": kind, "format": "date"})
        elif ptype in ("keyword", "enum"):
            out.append({**base, "id": f"distintos_{prop['id']}", "label": f"Distintos {prop.get('label') or prop['id']}", "kind": "distintos"})
    return out


def _dataset_from_object_type(obj_type: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    binding = obj_type.get("binding") or {}
    kind = binding.get("kind")
    props = _prop_map(obj_type)
    requires_session = bool(binding.get("scoped") or obj_type.get("requires_session"))

    if kind in ("es", "aggregation"):
        index = binding.get("index")
        if not index:
            return None
        dimensions = [entry for entry in (_field_entry(prop, obj_type) for prop in obj_type.get("properties") or []) if entry]
        measures = [{"id": "contagem", "label": "Contagem", "kind": "contagem", "format": "integer", "source": obj_type["id"]}]
        measures.extend(_measures_for(obj_type, props))

        primary: Optional[Dict[str, Any]] = None
        if kind == "aggregation":
            curated = CURATED_DIMENSIONS.get(obj_type["id"])
            prop = props.get(curated["prop"]) if curated and curated.get("prop") else None
            field = (prop or {}).get("field") or binding.get("field")
            if field:
                primary = {
                    "id": curated["id"] if curated else (binding.get("field") or "valor"),
                    "label": curated["label"] if curated else obj_type["label"],
                    "type": prop.get("type") if prop else "keyword",
                    "field": field,
                    "nested": (prop or {}).get("nested") or binding.get("nested"),
                    "unit": (prop or {}).get("unit"),
                    "values_from": (prop or {}).get("values_from"),
                    "sortable": True,
                    "filterable": True,
                    "groupable": True,
                    "pk": True,
                    "primary": True,
                    "from": obj_type["id"],
                }
        for entry in CURATED_MEASURES.get(obj_type["id"], []):
            # A propriedade pode pertencer ao tipo irmão do mesmo índice
            # (ex.: `valor` do CPV usa `preco` do contrato): o campo é resolvido
            # depois, em `_enrich_dataset`.
            target = props.get(entry["prop"]) if entry.get("prop") else None
            measures.append({
                "id": entry["id"],
                "label": entry["label"],
                "kind": entry["kind"],
                "field": (target or {}).get("field"),
                "nested": (target or {}).get("nested"),
                "unit": (target or {}).get("unit"),
                "format": FORMAT_BY_UNIT.get(str((target or {}).get("unit"))),
                "prop": entry.get("prop"),
                "source": obj_type["id"],
            })
        return {
            "id": obj_type["id"],
            "label": obj_type.get("plural") or obj_type["label"],
            "description": obj_type.get("description"),
            "domain": obj_type.get("domain"),
            "icon": obj_type.get("icon"),
            "kind": kind,
            "index": index,
            "requires_session": requires_session,
            "aggregation": True,
            "records": kind == "es",
            "primary_dimension": primary,
            "dimensions": dimensions,
            "measures": measures,
            "filter": binding.get("filter"),
            "search_fields": binding.get("search_fields") or [],
            "default_sort": binding.get("default_sort"),
            "binding": binding,
            "source": "ontologia",
        }

    if kind == "derived":
        return {
            "id": obj_type["id"],
            "label": obj_type.get("plural") or obj_type["label"],
            "description": obj_type.get("description"),
            "domain": obj_type.get("domain"),
            "icon": obj_type.get("icon"),
            "kind": "derived",
            "index": None,
            "requires_session": requires_session,
            "aggregation": False,
            "records": True,
            "primary_dimension": None,
            "dimensions": [],
            "measures": [],
            "filter": None,
            "search_fields": [],
            "default_sort": binding.get("default_sort"),
            "binding": binding,
            "source": "ontologia",
            "notes": [
                "Dataset resolvido em memória (combina várias fontes): disponível para listagens "
                "e detalhe, sem agregações para não apresentar amostras como totais."
            ],
        }
    return None


# --- Datasets locais (fontes que não vivem no Elasticsearch) ----------------
LOCAL_DATASETS: List[Dict[str, Any]] = [
    {
        "id": "documento",
        "label": "Documentos Office",
        "description": "Notas, relatórios, atas e páginas escritas na plataforma (Office).",
        "domain": "office",
        "icon": "book-open",
        "kind": "local",
        "index": None,
        "requires_session": False,
        "aggregation": True,
        "records": False,
        "primary_dimension": None,
        "dimensions": [
            {"id": "tipo", "label": "Tipo", "type": "keyword", "field": "kind_label"},
            {"id": "pasta", "label": "Pasta", "type": "keyword", "field": "folder"},
            {"id": "autor", "label": "Autor", "type": "keyword", "field": "author"},
            {"id": "etiqueta", "label": "Etiqueta", "type": "keyword", "field": "tags", "multi": True},
            {"id": "fixado", "label": "Fixado", "type": "boolean", "field": "pinned"},
            {"id": "criado_em", "label": "Criado em", "type": "date", "field": "created_at"},
            {"id": "atualizado_em", "label": "Atualizado em", "type": "date", "field": "updated_at"},
        ],
        "measures": [
            {"id": "contagem", "label": "Documentos", "kind": "contagem", "format": "integer"},
            {"id": "palavras", "label": "Palavras (soma)", "kind": "soma", "field": "words", "format": "integer"},
            {"id": "palavras_media", "label": "Palavras (média)", "kind": "media", "field": "words", "format": "number"},
            {"id": "palavras_maximo", "label": "Palavras (máximo)", "kind": "maximo", "field": "words", "format": "integer"},
        ],
        "notes": ["Carregado de `data/office/office.json`; agregação calculada em memória."],
        "source": "office",
    },
    {
        "id": "dossie",
        "label": "Dossiês 360",
        "description": "Dossiês temáticos guardados na Pesquisa 360, com volumetria e sentimento.",
        "domain": "search360",
        "icon": "telescope",
        "kind": "local",
        "index": None,
        "requires_session": False,
        "aggregation": True,
        "records": False,
        "primary_dimension": None,
        "dimensions": [
            {"id": "termo", "label": "Tema", "type": "keyword", "field": "term"},
            {"id": "projeto", "label": "Projeto", "type": "keyword", "field": "project_name"},
            {"id": "autor", "label": "Autor", "type": "keyword", "field": "author"},
            {"id": "etiqueta", "label": "Etiqueta", "type": "keyword", "field": "tags", "multi": True},
            {"id": "modo_sintese", "label": "Modo da síntese", "type": "keyword", "field": "synthesis_mode"},
            {"id": "sentimento", "label": "Sentimento", "type": "keyword", "field": "sentiment_label"},
            {"id": "criado_em", "label": "Criado em", "type": "date", "field": "created_at"},
            {"id": "atualizado_em", "label": "Atualizado em", "type": "date", "field": "updated_at"},
        ],
        "measures": [
            {"id": "contagem", "label": "Dossiês", "kind": "contagem", "format": "integer"},
            {"id": "itens", "label": "Itens recolhidos", "kind": "soma", "field": "items", "format": "integer"},
            {"id": "metricas", "label": "Métricas", "kind": "soma", "field": "metrics", "format": "integer"},
            {"id": "evidencias", "label": "Evidências citadas", "kind": "soma", "field": "evidence", "format": "integer"},
            {"id": "sentimento_medio", "label": "Sentimento médio", "kind": "media", "field": "sentiment_mean", "format": "number"},
        ],
        "notes": ["Carregado do registo da Ontologia (chaves `dossiers`); agregação calculada em memória."],
        "source": "search360",
    },
    {
        "id": "conta_email",
        "label": "Contas de Email",
        "description": "Contas de correio configuradas pelo utilizador, por fornecedor e estado de sincronização.",
        "domain": "email",
        "icon": "mail",
        "kind": "local",
        "index": None,
        "requires_session": True,
        "aggregation": True,
        "records": False,
        "primary_dimension": None,
        "dimensions": [
            {"id": "fornecedor", "label": "Fornecedor", "type": "keyword", "field": "provider"},
            {"id": "conta", "label": "Conta", "type": "keyword", "field": "email_address"},
            {"id": "predefinida", "label": "Predefinida", "type": "boolean", "field": "is_default"},
            {"id": "com_erro", "label": "Com erro", "type": "boolean", "field": "has_error"},
            {"id": "ultima_sincronizacao", "label": "Última sincronização", "type": "date", "field": "last_sync_at"},
        ],
        "measures": [{"id": "contagem", "label": "Contas", "kind": "contagem", "format": "integer"}],
        "notes": ["Só as contas do utilizador autenticado; as mensagens vivem no servidor IMAP e não são indexadas."],
        "source": "email",
    },
    {
        "id": "recolha",
        "label": "Recolha (scraping)",
        "description": "Itens extraídos pelos coletores: fonte, etiquetas e data de recolha.",
        "domain": "scraper",
        "icon": "database",
        "kind": "es",
        "index": "finance_scraped",
        "requires_session": False,
        "aggregation": True,
        "records": True,
        "primary_dimension": None,
        "count_field": "source_id",
        "dimensions": [
            {"id": "fonte", "label": "Fonte", "type": "keyword", "field": "source_name"},
            {"id": "fonte_id", "label": "ID da fonte", "type": "keyword", "field": "source_id"},
            {"id": "execucao", "label": "Execução", "type": "keyword", "field": "run_id"},
            {"id": "etiqueta", "label": "Etiqueta", "type": "keyword", "field": "tags"},
            {"id": "titulo", "label": "Título", "type": "text", "field": "title"},
            {"id": "acionador", "label": "Acionador", "type": "keyword", "field": "trigger"},
            {"id": "recolhido_em", "label": "Recolhido em", "type": "date", "field": "scraped_at"},
        ],
        "measures": [{"id": "contagem", "label": "Itens", "kind": "contagem", "format": "integer"}],
        "search_fields": ["title", "summary", "text"],
        "default_sort": {"field": "scraped_at", "order": "desc"},
        "binding": {"kind": "es", "index": "finance_scraped", "id_field": "item_id",
                    "search_fields": ["title", "summary", "text"],
                    "default_sort": {"field": "scraped_at", "order": "desc"}},
        "notes": ["Índice `finance_scraped`: cada documento é um item extraído por uma fonte."],
        "source": "builtin",
    },
]


def _local_rows(dataset_id: str, owner: Optional[str]) -> List[Dict[str, Any]]:
    """Carrega as linhas de um dataset local (fontes fora do Elasticsearch)."""
    if dataset_id == "documento":
        from api import office_store

        return office_store.list_documents(limit=MAX_LOCAL_ROWS)["items"]

    if dataset_id == "dossie":
        from api import search360_store

        listing = search360_store.list_dossiers(limit=5000)
        projects = {str(item.get("id")): item.get("name") for item in search360_store.list_projects().get("items", [])}
        rows = []
        for item in listing.get("items", []):
            rows.append({**item, "project_name": projects.get(str(item.get("project_id")))})
        return rows

    if dataset_id == "conta_email":
        if not owner:
            return []
        from api import email_service

        accounts = email_service.list_accounts(owner).get("items", [])
        return [
            {
                **account,
                "provider": account.get("provider"),
                "has_error": bool(account.get("last_error")),
            }
            for account in accounts
        ]
    return []


def _local_spec(dataset_id: str) -> Optional[Dict[str, Any]]:
    for spec in LOCAL_DATASETS:
        if spec["id"] == dataset_id:
            return spec
    return None


# ---------------------------------------------------------------------------
# Catálogo público
# ---------------------------------------------------------------------------
def _index_fields_map() -> Dict[str, List[Dict[str, Any]]]:
    """Campos por índice, a partir de todos os tipos da Ontologia.

    Permite que um dataset agregado (ex.: CPV) herde as dimensões e medidas do
    tipo que vive no mesmo índice (ex.: Contrato) — é assim que se analisam
    códigos CPV por ano, região ou procedimento.
    """
    out: Dict[str, List[Dict[str, Any]]] = {}
    for obj_type in registry.load_ontology().get("object_types", []):
        binding = obj_type.get("binding") or {}
        if binding.get("kind") not in ("es",):
            continue
        index = binding.get("index")
        if not index:
            continue
        entry = out.setdefault(index, [])
        entry.append(obj_type)
    return out


def _enrich_dataset(spec: Dict[str, Any], index_types: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    """Junta dimensões/medidas do tipo com as dos tipos irmãos no mesmo índice."""
    dataset = dict(spec)
    if dataset["kind"] == "aggregation" and dataset.get("primary_dimension"):
        dataset["dimensions"] = [dataset["primary_dimension"], *dataset.get("dimensions", [])]
    dataset.setdefault("count_field", None)
    sibling_measures: List[Dict[str, Any]] = []
    if dataset["kind"] in ("es", "aggregation") and dataset.get("index"):
        for sibling in index_types.get(dataset["index"], []):
            if sibling["id"] == dataset["id"]:
                continue
            sibling_spec = _dataset_from_object_type(sibling)
            if not sibling_spec:
                continue
            dataset["dimensions"].extend(sibling_spec["dimensions"])
            sibling_measures.extend(sibling_spec["measures"])
            dataset["measures"].extend(sibling_spec["measures"])
            if not dataset.get("count_field"):
                dataset["count_field"] = _guess_count_field(sibling)

    # Medidas curadas de datasets agregados (ex.: "valor" no CPV): a propriedade
    # pertence ao tipo irmão que vive no mesmo índice (ex.: `contrato.preco`).
    by_prop = {entry.get("prop"): entry for entry in sibling_measures if entry.get("prop") and entry.get("field")}
    for entry in dataset["measures"]:
        if entry.get("field") or entry["kind"] == "contagem":
            continue
        source = by_prop.get(entry.get("prop"))
        if source:
            entry["field"] = source["field"]
            entry["nested"] = source.get("nested")
            entry["unit"] = entry.get("unit") or source.get("unit")
            entry["format"] = entry.get("format") or source.get("format")

    # Deduplicar por id (o próprio tipo ganha).
    seen_dims: set = set()
    deduped_dims: List[Dict[str, Any]] = []
    for entry in dataset["dimensions"]:
        if entry["id"] in seen_dims or not entry.get("groupable", True):
            continue
        seen_dims.add(entry["id"])
        deduped_dims.append(entry)
    seen_measures: set = set()
    deduped_measures = []
    for entry in dataset["measures"]:
        if entry["id"] in seen_measures or not entry.get("field") and entry["kind"] != "contagem":
            continue
        seen_measures.add(entry["id"])
        deduped_measures.append(entry)
    dataset["dimensions"] = deduped_dims
    dataset["measures"] = deduped_measures
    if not dataset.get("count_field"):
        dataset["count_field"] = _guess_count_field_from_spec(dataset)
    dataset["filters"] = [entry for entry in dataset["dimensions"] if entry.get("filterable")]
    dataset["defaults"] = _defaults_for(dataset)
    dataset["suggestions"] = _suggestions_for(dataset)
    return dataset


def _guess_count_field(obj_type: Dict[str, Any]) -> Optional[str]:
    """Campo do índice presente em todos os documentos (para contar em contexto aninhado)."""
    binding = obj_type.get("binding") or {}
    candidates: List[str] = []
    id_field = binding.get("id_field")
    if id_field and not str(id_field).startswith("_") and "." not in str(id_field):
        candidates.append(id_field)
    for prop in obj_type.get("properties") or []:
        field = prop.get("field")
        if not field or "." in str(field) or prop.get("type") == "text":
            continue
        candidates.append(str(field))
    return candidates[0] if candidates else None


def _guess_count_field_from_spec(dataset: Dict[str, Any]) -> Optional[str]:
    if dataset.get("count_field"):
        return dataset["count_field"]
    binding = dataset.get("binding") or {}
    id_field = binding.get("id_field")
    if id_field and not str(id_field).startswith("_") and "." not in str(id_field):
        return str(id_field)
    for entry in dataset.get("dimensions", []):
        if entry.get("field") and "." not in str(entry["field"]) and entry.get("type") != "text":
            return entry["field"]
    for entry in dataset.get("measures", []):
        if entry.get("field") and "." not in str(entry["field"]):
            return entry["field"]
    return None


def _defaults_for(dataset: Dict[str, Any]) -> Dict[str, Any]:
    """Escolhe uma dimensão e uma medida de partida sensatas."""
    dims = dataset.get("dimensions", [])
    measures = dataset.get("measures", [])
    date_dims = [d for d in dims if d.get("type") == "date"]
    year_dims = [d for d in dims if d["id"] in ("ano", "year")]
    dimension = (year_dims or [d for d in dims if d["id"] in ("regiao", "ticker", "fonte")] or date_dims or dims)
    value_dims = [d for d in dims if d.get("type") in ("keyword", "enum")]
    measure_ids = [m["id"] for m in measures]
    value_measure = next((m for m in measure_ids if m in ("valor", "soma_preco", "soma_valor", "soma_amount")), None)
    return {
        "dimension": (dimension[0]["id"] if dimension else None),
        "dimension_interval": "ano" if (dimension and dimension[0].get("type") == "date") else None,
        "measure": value_measure or (measure_ids[1] if len(measure_ids) > 1 else "contagem"),
        "table_dimension": (value_dims[0]["id"] if value_dims else (dimension[0]["id"] if dimension else None)),
        "limit": DEFAULT_LIMIT,
    }


def _suggestions_for(dataset: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Visuais sugeridos (atalhos de análise) para o dataset."""
    out: List[Dict[str, Any]] = []
    dims = {entry["id"]: entry for entry in dataset.get("dimensions", [])}
    measures = {entry["id"]: entry for entry in dataset.get("measures", [])}
    value = next((mid for mid in ("valor", "soma_preco", "soma_fecho", "soma_amount", "soma_valor") if mid in measures), "contagem")
    primary = dataset.get("primary_dimension", {}).get("id") if dataset.get("primary_dimension") else None

    def add(title: str, dimension: Optional[str], measure: str, chart: str, interval: Optional[str] = None) -> None:
        if not dimension or dimension not in dims:
            return
        entry = dims[dimension]
        out.append({
            "title": title,
            "dimension": dimension,
            "dimension_label": entry.get("label"),
            "interval": interval,
            "measure": measure,
            "measure_label": measures.get(measure, {}).get("label", measure),
            "chart": chart,
        })

    if "ano" in dims and value in measures:
        add("Evolução por ano", "ano", value, "area")
    if "regiao" in dims and value in measures:
        add("Valor por região", "regiao", value, "bar")
    if primary:
        add("Top 10", primary, value, "bar-h", interval=None)
    if "cpv" in dims and value in measures:
        add("Top CPV por valor", "cpv", value, "bar-h")
    if "tipo_contrato" in dims:
        add("Tipo de contrato", "tipo_contrato", "contagem", "donut")
    if "procedimento" in dims:
        add("Procedimento", "procedimento", "contagem", "donut")
    if "data" in dims:
        add("Série temporal", "data", value, "line", interval="mes")
    if "publicado" in dims:
        add("Notícias por mês", "publicado", "contagem", "line", interval="mes")
    if "ticker" in dims:
        add("Por instrumento", "ticker", "contagem", "bar-h")
    if "pasta" in dims:
        add("Documentos por pasta", "pasta", "contagem", "bar")
    return out[:8]


def _all_datasets(scope: Optional[Dict[str, Any]] = None, client: Any = None) -> List[Dict[str, Any]]:
    index_types = _index_fields_map()
    datasets: List[Dict[str, Any]] = []
    for obj_type in registry.load_ontology().get("object_types", []):
        spec = _dataset_from_object_type(obj_type)
        if not spec:
            continue
        datasets.append(_enrich_dataset(spec, index_types))
    for spec in LOCAL_DATASETS:
        datasets.append(_enrich_dataset(dict(spec), index_types))
    return datasets


def get_dataset(dataset_id: str) -> Dict[str, Any]:
    """Devolve a definição completa de um dataset (sem verificar sessão)."""
    for dataset in _all_datasets():
        if dataset["id"] == dataset_id:
            dataset["available"] = True
            return dataset
    raise KeyError(f"Dataset desconhecido: {dataset_id}")


def _dataset_summary(dataset: Dict[str, Any], scope: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    available = True
    note = None
    if dataset.get("requires_session") and not scope:
        available = False
        note = "Requer sessão iniciada: os dados são privados do utilizador."
    return {
        "id": dataset["id"],
        "label": dataset["label"],
        "description": dataset.get("description"),
        "domain": dataset.get("domain"),
        "icon": dataset.get("icon"),
        "kind": dataset["kind"],
        "index": dataset.get("index"),
        "requires_session": bool(dataset.get("requires_session")),
        "aggregation": bool(dataset.get("aggregation")),
        "records": bool(dataset.get("records")),
        "dimensions": len(dataset.get("dimensions") or []),
        "measures": len(dataset.get("measures") or []),
        "available": available,
        "note": note,
        "suggestions": dataset.get("suggestions") or [],
        "defaults": dataset.get("defaults") or {},
        "source": dataset.get("source"),
        "notes": dataset.get("notes") or [],
    }


def catalog(scope: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Catálogo de datasets para a UI (agrupado por domínio)."""
    datasets = [_dataset_summary(dataset, scope) for dataset in _all_datasets()]
    domains = []
    for domain in registry.load_ontology().get("domains", []) + LOCAL_DOMAINS:
        members = [dataset for dataset in datasets if dataset.get("domain") == domain["id"]]
        if not members:
            continue
        domains.append({**domain, "datasets": [member["id"] for member in members]})
    unknown = [dataset for dataset in datasets if not any(domain["id"] == dataset.get("domain") for domain in domains)]
    if unknown:
        domains.append({
            "id": "local",
            "label": "Plataforma",
            "description": "Dados produzidos dentro da plataforma (Office, dossiês, Email, recolha).",
            "accent": "148,163,184",
            "datasets": [dataset["id"] for dataset in unknown],
        })
    return {
        "datasets": datasets,
        "domains": domains,
        "measure_kinds": [
            {"id": key, "label": value["label"], "needs_field": value["field"], "additive": value["additive"]}
            for key, value in MEASURE_KINDS.items()
        ],
        "intervals": [{"id": key, "label": key.capitalize()} for key in INTERVALS],
        "chart_types": [
            {"id": "bar", "label": "Barras", "dimensions": 1},
            {"id": "bar-h", "label": "Barras horizontais", "dimensions": 1},
            {"id": "line", "label": "Linhas", "dimensions": 1, "time": True},
            {"id": "area", "label": "Área", "dimensions": 1, "time": True},
            {"id": "donut", "label": "Circular / Donut", "dimensions": 1},
            {"id": "treemap", "label": "Treemap", "dimensions": 2},
            {"id": "scatter", "label": "Dispersão", "dimensions": 1, "measures": 2},
            {"id": "stacked", "label": "Barras empilhadas", "dimensions": 2},
            {"id": "matrix", "label": "Tabela / matriz", "dimensions": 3},
            {"id": "kpi", "label": "Indicador (KPI)", "dimensions": 0},
        ],
        "limits": {
            "max_dimensions": MAX_DIMENSIONS,
            "max_measures": MAX_MEASURES,
            "max_formulas": MAX_FORMULAS,
            "max_rows": MAX_ROWS,
            "default_limit": DEFAULT_LIMIT,
            "notes": [
                "As agregações são calculadas no Elasticsearch (nunca extrapoladas de amostras).",
                "Filtros de texto usam `match`; campos de texto são agregados por runtime field quando o mapping não tem `keyword`.",
                "Nos datasets locais (Office, dossiês, Email) a agregação é calculada em memória sobre o ficheiro da plataforma.",
            ],
        },
        "ontology": registry.active_ontology_id(),
        "generated_at": _now_iso(),
    }


LOCAL_DOMAINS: List[Dict[str, str]] = [
    {"id": "office", "label": "Office", "description": "Documentos escritos na plataforma.", "accent": "14,165,233"},
    {"id": "search360", "label": "Pesquisa 360", "description": "Dossiês e projetos temáticos.", "accent": "99,102,241"},
    {"id": "email", "label": "Email", "description": "Contas de correio configuradas.", "accent": "20,184,166"},
    {"id": "scraper", "label": "Recolha", "description": "Itens extraídos pelos coletores.", "accent": "139,92,246"},
]


# ---------------------------------------------------------------------------
# Resolução das medidas/dimensões pedidas
# ---------------------------------------------------------------------------
def _resolve_dimensions(dataset: Dict[str, Any], raw: Sequence[Any]) -> Tuple[List[Dict[str, Any]], List[str]]:
    props = {entry["id"]: entry for entry in dataset.get("dimensions") or []}
    out: List[Dict[str, Any]] = []
    problems: List[str] = []
    for item in list(raw or [])[:MAX_DIMENSIONS]:
        if isinstance(item, str):
            item = {"id": item}
        if not isinstance(item, dict):
            continue
        spec = props.get(str(item.get("id")))
        if not spec:
            problems.append(f"Dimensão desconhecida: {item.get('id')}")
            continue
        entry = dict(spec)
        interval = str(item.get("interval") or "").lower()
        if entry.get("type") == "date":
            entry["interval"] = interval if interval in INTERVALS else "mes"
        out.append(entry)
    return out, problems


def _resolve_measures(dataset: Dict[str, Any], raw: Sequence[Any]) -> Tuple[List[Dict[str, Any]], List[str]]:
    catalog_map = {entry["id"]: entry for entry in dataset.get("measures") or []}
    props = {entry["id"]: entry for entry in dataset.get("dimensions") or []}
    out: List[Dict[str, Any]] = []
    problems: List[str] = []
    for item in list(raw or [])[:MAX_MEASURES]:
        if isinstance(item, str):
            item = {"id": item}
        if not isinstance(item, dict):
            continue
        kind = str(item.get("kind") or "").lower()
        measure_id = str(item.get("id") or "")
        if not kind and measure_id in catalog_map:
            entry = dict(catalog_map[measure_id])
        elif not kind and not item.get("field"):
            problems.append(f"Medida desconhecida neste dataset: {measure_id or item}")
            continue
        else:
            if kind and kind not in MEASURE_KINDS:
                problems.append(f"Tipo de medida desconhecido: {kind}")
                continue
            field_id = item.get("field")
            prop = props.get(str(field_id)) if field_id else None
            if kind != "contagem" and not prop:
                if measure_id in catalog_map:
                    entry = dict(catalog_map[measure_id])
                else:
                    problems.append(f"Medida sem campo válido: {measure_id or kind}")
                    continue
            else:
                entry = {
                    "id": measure_id or f"{kind}_{field_id}",
                    "label": item.get("label") or (f"{MEASURE_KINDS[kind]['label']} de {prop.get('label')}" if prop else MEASURE_KINDS[kind]["label"]),
                    "kind": kind or "contagem",
                    "field": (prop or {}).get("field"),
                    "nested": (prop or {}).get("nested"),
                    "unit": (prop or {}).get("unit"),
                    "format": item.get("format") or ("integer" if (prop or {}).get("type") == "number" and not (prop or {}).get("unit") else None),
                    "prop": (prop or {}).get("id"),
                }
                if kind == "contagem":
                    entry["field"] = None
        entry["kind"] = entry.get("kind") or "contagem"
        if entry["kind"] != "contagem" and not entry.get("field"):
            entry["field"] = _field_from_catalog(dataset, entry.get("prop"))
        if entry["kind"] != "contagem" and not entry.get("field"):
            problems.append(f"A medida «{entry.get('label')}» não tem campo associado neste dataset.")
            continue
        out.append(entry)
    if not out:
        out = [{"id": "contagem", "label": "Contagem", "kind": "contagem", "format": "integer"}]
    return out, problems


def _field_from_catalog(dataset: Dict[str, Any], prop_id: Optional[str]) -> Optional[str]:
    if not prop_id:
        return None
    if dataset.get("kind") == "local":
        entry = next((item for item in dataset.get("dimensions") or [] if item["id"] == prop_id), None)
        return (entry or {}).get("field")
    return None


def _resolve_formulas(raw: Sequence[Any]) -> Tuple[List[Dict[str, Any]], List[str]]:
    out: List[Dict[str, Any]] = []
    problems: List[str] = []
    for item in list(raw or [])[:MAX_FORMULAS]:
        if not isinstance(item, dict):
            continue
        expression = str(item.get("expression") or "").strip()
        if not expression:
            continue
        out.append({
            "id": str(item.get("id") or f"f{len(out)}"),
            "label": str(item.get("label") or f"Medida {len(out) + 1}"),
            "expression": expression,
            "format": item.get("format") or "number",
            "unit": item.get("unit"),
        })
    return out, problems


def _nested_of(entry: Dict[str, Any]) -> Optional[str]:
    return entry.get("nested") or None


def _index_of_dataset(dataset: Dict[str, Any]) -> Optional[str]:
    return dataset.get("index")


# ---------------------------------------------------------------------------
# Construção das agregações do Elasticsearch
# ---------------------------------------------------------------------------
def _wrap_dimension(index: int, target_nested: Optional[str], ctx_nested: Optional[str], node: Dict[str, Any]) -> Dict[str, Any]:
    """Coloca a agregação da dimensão no contexto aninhado certo."""
    name = f"d{index}"
    if target_nested == ctx_nested:
        return {name: node}
    if target_nested is None:
        return {f"rn{index}": {"reverse_nested": {}, "aggs": {name: node}}}
    if ctx_nested is None:
        return {f"nd{index}": {"nested": {"path": target_nested}, "aggs": {name: node}}}
    return {
        f"rn{index}": {
            "reverse_nested": {},
            "aggs": {f"nd{index}": {"nested": {"path": target_nested}, "aggs": {name: node}}},
        }
    }


def _read_dimension_container(container: Dict[str, Any], index: int, target_nested: Optional[str], ctx_nested: Optional[str]) -> Dict[str, Any]:
    if target_nested == ctx_nested:
        return container.get(f"d{index}") or {}
    if target_nested is None:
        return (container.get(f"rn{index}") or {}).get(f"d{index}") or {}
    if ctx_nested is None:
        return (container.get(f"nd{index}") or {}).get(f"d{index}") or {}
    return (((container.get(f"rn{index}") or {}).get(f"nd{index}") or {}).get(f"d{index}")) or {}


def _metric_node(kind: str, target_field: str, wrap_needed: bool) -> Dict[str, Any]:
    if kind == "soma":
        body: Dict[str, Any] = {"sum": {"field": target_field}}
    elif kind == "media":
        body = {"avg": {"field": target_field}}
    elif kind == "minimo":
        body = {"min": {"field": target_field}}
    elif kind == "maximo":
        body = {"max": {"field": target_field}}
    elif kind == "distintos":
        body = {"cardinality": {"field": target_field, "precision_threshold": 40000}}
    else:  # mediana / p90
        body = {"percentiles": {"field": target_field, "percents": [50 if kind == "mediana" else 90], "keyed": True}}
    if wrap_needed:
        # A medida vive fora do objeto aninhado da dimensão: volta-se ao documento
        # original (`reverse_nested`) para somar valores do contrato, não da parte.
        return {"reverse_nested": {}, "aggs": {"n": body}}
    return body


def _build_metrics(
    dataset: Dict[str, Any],
    measures: Sequence[Dict[str, Any]],
    ctx_nested: Optional[str],
    client: Any,
    index: Optional[str],
    runtime: Dict[str, Any],
    properties: Dict[str, Any],
    prefix: str = "",
) -> Tuple[Dict[str, Any], List[List[str]], List[Dict[str, Any]], Dict[str, str]]:
    """Agregações das medidas + caminho de leitura de cada valor.

    Devolve também o **caminho de ordenação** de cada medida dentro dos buckets
    (`m0` ou `m0>n`, quando a medida é embrulhada em `reverse_nested`), necessário
    para pedir ao Elasticsearch os grupos certos (top-N por valor, não por
    frequência).

    `prefix` distingue as agregações de totais (`t_`) das agregações por grupo:
    cada nome na árvore `aggs` do Elasticsearch só pode ter **um** tipo de
    agregação, pelo que os totais não podem ser um bloco com vários nomes.
    """
    aggs: Dict[str, Any] = {}
    reads: List[List[str]] = []
    resolution: List[Dict[str, Any]] = []
    order_paths: Dict[str, str] = {}
    count_field = dataset.get("count_field")
    for position, measure in enumerate(measures):
        name = f"{prefix}m{position}"
        kind = measure["kind"]
        if kind == "contagem":
            if ctx_nested:
                if count_field:
                    aggs[name] = {"reverse_nested": {}, "aggs": {"n": {"value_count": {"field": count_field}}}}
                    reads.append([name, "n", "value"])
                    order_paths[name] = f"{name}>n"
                else:
                    aggs[name] = {"reverse_nested": {}}
                    reads.append([name, "doc_count"])
            else:
                reads.append([])
                order_paths[name] = f"{name}"
            resolution.append({"id": measure["id"], "label": measure.get("label"), "kind": kind, "field": None})
            continue
        field = measure.get("field")
        target = _agg_target(client, index, field, properties) if field else {"field": None, "runtime": None}
        if target.get("runtime"):
            runtime.update(target["runtime"])
        field_nested = _nested_of(measure)
        wrap = bool(ctx_nested) and not (field and str(field).startswith(f"{ctx_nested}.")) and field_nested != ctx_nested
        aggs[name] = _metric_node(kind, target["field"], wrap)
        if kind in ("mediana", "p90"):
            reads.append([name, "values", "50.0" if kind == "mediana" else "90.0"])
        elif wrap:
            reads.append([name, "n", "value"])
            order_paths[name] = f"{name}>n"
        else:
            reads.append([name, "value"])
            order_paths[name] = f"{name}"
        resolution.append({
            "id": measure["id"],
            "label": measure.get("label"),
            "kind": kind,
            "field": field,
            "agg_field": target["field"],
            "note": target.get("note"),
        })
    return aggs, reads, resolution, order_paths


def _read_metrics(container: Dict[str, Any], reads: Sequence[List[str]], doc_count: int) -> List[Any]:
    values: List[Any] = []
    for path in reads:
        if not path:
            values.append(doc_count)
            continue
        node: Any = container
        for key in path:
            if isinstance(node, dict) and key in node:
                node = node[key]
            else:
                node = None
                break
        values.append(node if isinstance(node, (int, float)) else None)
    return values


def _build_dimension_chain(
    dataset: Dict[str, Any],
    dimensions: Sequence[Dict[str, Any]],
    measures: Sequence[Dict[str, Any]],
    client: Any,
    index: Optional[str],
    runtime: Dict[str, Any],
    properties: Dict[str, Any],
    limit: int,
    inner_limit: int,
    order_metric: Optional[str],
    missing_label: Optional[str],
) -> Tuple[Dict[str, Any], List[List[str]], List[Dict[str, Any]]]:

    def build_level(position: int, ctx_nested: Optional[str]) -> Tuple[Dict[str, Any], List[List[str]], List[Dict[str, Any]], Dict[str, str]]:
        if position >= len(dimensions):
            return _build_metrics(dataset, measures, ctx_nested, client, index, runtime, properties)
        dimension = dimensions[position]
        child_aggs, reads, resolution, order_paths = build_level(position + 1, _nested_of(dimension))
        target = _agg_target(client, index, dimension["field"], properties)
        if target.get("runtime"):
            runtime.update(target["runtime"])
        resolution.insert(0, {
            "id": dimension["id"],
            "label": dimension.get("label"),
            "type": dimension.get("type"),
            "field": dimension["field"],
            "agg_field": target["field"],
            "nested": _nested_of(dimension),
            "interval": dimension.get("interval"),
            "unit": dimension.get("unit"),
        })
        if dimension.get("type") == "date":
            node: Dict[str, Any] = {
                "date_histogram": {
                    "field": target["field"],
                    "calendar_interval": INTERVALS.get(str(dimension.get("interval") or "mes"), "month"),
                    "min_doc_count": 1,
                    "order": {"_key": "asc"},
                },
                "aggs": child_aggs,
            }
        else:
            size = limit if position == 0 else inner_limit
            terms: Dict[str, Any] = {"field": target["field"], "size": size}
            path = order_paths.get(order_metric or "") if position == 0 else None
            if path:
                terms["order"] = {path: "desc"}
            elif position == 0 and _is_temporal(dimension):
                # Série temporal: por ordem de chave (cronológica) e com todos os
                # períodos, senão faltariam anos/meses menos movimentados.
                terms["order"] = {"_key": "asc"}
            else:
                terms["order"] = {"_count": "desc"}
            if missing_label and position == 0 and order_paths and target.get("type") in ("keyword", "boolean", "flattened", "ip"):
                terms["missing"] = missing_label
            node = {"terms": terms, "aggs": child_aggs}
        return _wrap_dimension(position, _nested_of(dimension), ctx_nested, node), reads, resolution, order_paths

    return build_level(0, None)


# ---------------------------------------------------------------------------
# Consulta principal
# ---------------------------------------------------------------------------
def _base_query(dataset: Dict[str, Any], filters: Dict[str, Any], search: Optional[str]) -> Tuple[Dict[str, Any], List[str]]:
    obj_type_id = dataset.get("id")
    problems: List[str] = []
    clauses: List[Dict[str, Any]] = []
    try:
        obj_type = ontology.get_object_type(obj_type_id)
    except Exception:
        obj_type = None
    if obj_type:
        filter_clauses, unknown = ontology._build_filters(obj_type, filters or {})  # noqa: SLF001 (reutilização deliberada)
        clauses.extend(filter_clauses)
        if unknown:
            problems.append("Filtros ignorados (não existem neste dataset): " + ", ".join(sorted(unknown)))
    elif dataset.get("kind") == "local":
        return {}, problems
    else:
        clauses.extend(_generic_filters(dataset, filters or {}))

    if dataset.get("filter"):
        clauses.append(dataset["filter"])

    query: Dict[str, Any] = {"bool": {}}
    if clauses:
        query["bool"]["filter"] = clauses
    search_fields = dataset.get("search_fields") or []
    if isinstance(search, str) and search.strip() and search_fields:
        query["bool"]["must"] = {
            "multi_match": {"query": search.strip(), "fields": search_fields, "operator": "and"}
        }
    if not query["bool"]:
        return {"match_all": {}}, problems
    return query, problems


def _generic_filters(dataset: Dict[str, Any], filters: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Filtros diretos para datasets sem tipo na Ontologia (built-in)."""
    props = {entry["id"]: entry for entry in dataset.get("dimensions") or []}
    clauses: List[Dict[str, Any]] = []
    for key, value in (filters or {}).items():
        entry = props.get(key)
        if not entry or value in (None, ""):
            continue
        field = entry["field"]
        ptype = entry.get("type")
        if ptype in ("number", "date"):
            bounds: Dict[str, Any] = {}
            if isinstance(value, dict):
                for src, dst in (("min", "gte"), ("gte", "gte"), ("max", "lte"), ("lte", "lte")):
                    if value.get(src) is not None:
                        bounds[dst] = value[src]
            elif isinstance(value, list) and len(value) == 2:
                bounds = {"gte": value[0], "lte": value[1]}
            else:
                bounds = {"gte": value, "lte": value}
            if bounds:
                clauses.append({"range": {field: bounds}})
        elif ptype in ("keyword", "enum"):
            values = [str(item) for item in _as_list(value)]
            clauses.append({"terms": {field: values}} if len(values) != 1 else {"term": {field: values[0]}})
        else:
            clauses.append({"match": {field: value}})
    return clauses


def _scope_for(scope: Optional[Dict[str, Any]], dataset: Dict[str, Any]) -> List[Dict[str, Any]]:
    if not dataset.get("requires_session"):
        return []
    if not scope:
        raise PermissionError(f"O dataset «{dataset['label']}» exige sessão iniciada.")
    if scope.get("see_all"):
        return []
    return [{"term": {"owner_id": scope.get("user_id")}}]


def run_query(
    payload: Dict[str, Any],
    scope: Optional[Dict[str, Any]] = None,
    es: Any = None,
) -> Dict[str, Any]:
    """Executa uma consulta analítica e devolve linhas prontas a desenhar."""
    started = time.time()
    dataset_id = str(payload.get("dataset") or "").strip()
    if not dataset_id:
        return {"error": "Escolha um dataset.", "rows": []}
    try:
        dataset = get_dataset(dataset_id)
    except KeyError as exc:
        return {"error": str(exc), "rows": []}

    notes: List[str] = list(dataset.get("notes") or [])
    dimensions, problems = _resolve_dimensions(dataset, payload.get("dimensions") or [])
    measures, measure_problems = _resolve_measures(dataset, payload.get("measures") or [])
    formulas, _ = _resolve_formulas(payload.get("formulas") or [])
    notes.extend(problems)
    notes.extend(measure_problems)

    limit = max(1, min(int(payload.get("limit") or DEFAULT_LIMIT), TERMS_MAX))
    inner_limit = max(1, min(int(payload.get("inner_limit") or min(limit, 20)), 200))
    top_n = payload.get("top_n")
    top_n = max(0, int(top_n)) if top_n not in (None, "") else None
    others = bool(payload.get("others"))
    missing_label = payload.get("missing_label")
    missing_label = "Não especificado" if missing_label is None else (str(missing_label) or None)
    sort_request = payload.get("sort") or {}
    if isinstance(sort_request, str):
        sort_request = {"by": sort_request}
    sort_by = str(sort_request.get("by") or "")
    sort_order = "asc" if str(sort_request.get("order") or "desc").lower() == "asc" else "desc"
    filters = payload.get("filters") if isinstance(payload.get("filters"), dict) else {}
    search = payload.get("search")

    if not dataset.get("aggregation"):
        return {
            "dataset": _dataset_summary(dataset, scope),
            "dimensions": dimensions,
            "measures": measures,
            "formulas": formulas,
            "columns": [],
            "rows": [],
            "totals": {},
            "meta": {"notes": notes + ["Este dataset não suporta agregações: use a tabela de registos."], "truncated": False},
            "error": None,
        }

    if dataset["kind"] == "local":
        return _run_local_query(dataset, payload, dimensions, measures, formulas, filters, search,
                                limit, top_n, others, sort_by, sort_order, scope, notes, started)

    client = es or get_es_client()
    if not client:
        return {"error": "Elasticsearch indisponível.", "rows": [], "meta": {"notes": notes}, "dataset": _dataset_summary(dataset, scope)}

    index = _index_of_dataset(dataset)
    properties = _index_properties(client, index or "")
    runtime: Dict[str, Any] = {}

    try:
        query, query_problems = _base_query(dataset, filters, search)
    except PermissionError as exc:
        return {"error": str(exc), "rows": [], "meta": {"notes": notes}}
    notes.extend(query_problems)
    try:
        scope_clauses = _scope_for(scope, dataset)
    except PermissionError as exc:
        return {"error": str(exc), "rows": [], "meta": {"notes": notes}}
    if scope_clauses:
        bool_query = query.setdefault("bool", {})
        bool_query.setdefault("filter", [])
        bool_query["filter"].extend(scope_clauses)

    # Dimensões temporais trazem a série completa (a UI limita com o Top-N).
    effective_limit = max(limit, 300) if (dimensions and _is_temporal(dimensions[0])) else limit

    order_metric = None
    if sort_by and dimensions:
        candidate = next((f"m{position}" for position, measure in enumerate(measures) if measure["id"] == sort_by), None)
        if candidate and sort_order == "desc" and MEASURE_KINDS[measures[int(candidate[1:])]["kind"]].get("additive"):
            order_metric = candidate

    aggs, reads, resolution, _order_paths = _build_dimension_chain(
        dataset, dimensions, measures, client, index, runtime, properties,
        effective_limit, inner_limit, order_metric, missing_label,
    )
    totals_aggs, totals_reads, _, _ = _build_metrics(dataset, measures, None, client, index, runtime, properties, prefix="t_")
    aggs.update(totals_aggs)

    body: Dict[str, Any] = {"size": 0, "query": query, "aggs": aggs, "track_total_hits": True}
    if runtime:
        body["runtime_mappings"] = runtime

    try:
        response = client.search(index=index, body=body, request_timeout=AGG_TIMEOUT_SECONDS)
    except Exception as exc:
        logger.warning("Consulta do Visualizador falhou (%s): %s", dataset_id, exc)
        return {"error": f"A consulta falhou: {exc}", "rows": [], "meta": {"notes": notes}, "dataset": _dataset_summary(dataset, scope)}

    total_hits = ((response.get("hits") or {}).get("total") or {}).get("value", 0)
    aggregations = response.get("aggregations") or {}
    rows, truncated = _flatten_buckets(dataset, dimensions, measures, reads, resolution, aggregations, total_hits, effective_limit, inner_limit)

    totals_values = _read_metrics(aggregations, totals_reads, total_hits)
    totals = {measure["id"]: totals_values[position] for position, measure in enumerate(measures)}

    notes.append(
        f"{total_hits:,} documento(s) correspondem aos filtros".replace(",", " ")
        + (f" (limite de {limit} por dimensão)." if dimensions else ".")
    )
    if truncated:
        notes.append("A lista de valores foi truncada pelo limite de grupos: aumente o limite ou aplique filtros.")
    if any(entry.get("note") for entry in resolution):
        for entry in resolution:
            if entry.get("note"):
                notes.append(f"{entry.get('label') or entry['id']}: {entry['note']}")

    result = _finalize_rows(rows, dimensions, measures, formulas, resolution, totals, sort_by, sort_order, top_n, others)
    result.update({
        "dataset": _dataset_summary(dataset, scope),
        "totals": totals,
        "meta": {
            "documents": total_hits,
            "elapsed_ms": round((time.time() - started) * 1000, 1),
            "truncated": truncated or result.get("meta", {}).get("truncated", False),
            "notes": notes + list(result.get("meta", {}).get("notes") or []),
            "index": index,
            "resolution": resolution,
        },
        "error": None,
    })
    return result


def _flatten_buckets(
    dataset: Dict[str, Any],
    dimensions: Sequence[Dict[str, Any]],
    measures: Sequence[Dict[str, Any]],
    reads: Sequence[List[str]],
    resolution: Sequence[Dict[str, Any]],
    aggregations: Dict[str, Any],
    total_hits: int,
    limit: int = DEFAULT_LIMIT,
    inner_limit: int = DEFAULT_LIMIT,
) -> Tuple[List[Dict[str, Any]], bool]:
    """Percorre as agregações e produz as linhas (produto das dimensões)."""
    rows: List[Dict[str, Any]] = []
    truncated = False

    if not dimensions:
        values = _read_metrics(aggregations, reads, total_hits)
        row = {measure["id"]: values[position] for position, measure in enumerate(measures)}
        row["_key"] = "total"
        return [row], False

    def walk(container: Dict[str, Any], position: int, prefix: Dict[str, Any], ctx_nested: Optional[str], series: List[Dict[str, Any]]) -> None:
        nonlocal truncated
        dimension = dimensions[position]
        node = _read_dimension_container(container, position, _nested_of(dimension), ctx_nested)
        buckets = node.get("buckets") or []
        if len(buckets) >= (limit if position == 0 else inner_limit):
            truncated = True
        for bucket in buckets:
            key = bucket.get("key_as_string") if bucket.get("key_as_string") is not None else bucket.get("key")
            label = bucket.get("key_as_string") or bucket.get("key")
            entry = {**prefix, dimension["id"]: key}
            labels = {**prefix.get("_labels", {}), dimension["id"]: label}
            entry["_labels"] = labels
            if position + 1 < len(dimensions):
                walk(bucket, position + 1, entry, _nested_of(dimension), series)
                continue
            values = _read_metrics(bucket, reads, bucket.get("doc_count", 0))
            row = dict(entry)
            for index, measure in enumerate(measures):
                row[measure["id"]] = values[index]
            row["_series"] = series + [{k: v for k, v in entry.items() if k in {d["id"] for d in dimensions}}]
            row["_key"] = "|".join(str(row.get(d["id"])) for d in dimensions)
            rows.append(row)

    walk(aggregations, 0, {}, None, [])
    return rows, truncated


# ---------------------------------------------------------------------------
# Datasets locais (agregação em memória)
# ---------------------------------------------------------------------------
def _local_value(row: Dict[str, Any], entry: Dict[str, Any]) -> Any:
    field = entry.get("field") or entry["id"]
    value: Any = row
    for part in str(field).split("."):
        if isinstance(value, dict):
            value = value.get(part)
        else:
            value = None
            break
    if isinstance(value, list):
        return [item for item in value if item is not None]
    return value


def _local_measure(rows: Sequence[Dict[str, Any]], measure: Dict[str, Any]) -> Optional[float]:
    kind = measure["kind"]
    if kind == "contagem":
        return float(len(rows))
    values: List[float] = []
    for row in rows:
        raw = _local_value(row, measure) if measure.get("field") else None
        for item in (raw if isinstance(raw, list) else [raw]):
            number = item if isinstance(item, (int, float)) and not isinstance(item, bool) else _to_number(item)
            if number is not None:
                values.append(float(number))
    if not values:
        return None
    if kind == "soma":
        return float(sum(values))
    if kind == "media":
        return float(sum(values) / len(values))
    if kind == "minimo":
        return float(min(values))
    if kind == "maximo":
        return float(max(values))
    if kind == "distintos":
        return float(len(set(values)))
    ordered = sorted(values)
    if kind == "mediana":
        middle = len(ordered) // 2
        return float(ordered[middle]) if len(ordered) % 2 else float((ordered[middle - 1] + ordered[middle]) / 2)
    if kind == "p90":
        index = max(0, min(len(ordered) - 1, int(round(0.9 * (len(ordered) - 1)))))
        return float(ordered[index])
    return None


def _local_bucket(value: Any, entry: Dict[str, Any]) -> Any:
    if entry.get("type") == "date" and isinstance(value, str):
        interval = INTERVALS.get(str(entry.get("interval") or "mes"), "month")
        text = value[:10]
        if interval == "year":
            return text[:4]
        if interval == "month":
            return text[:7]
        if interval == "quarter":
            month = int(text[5:7] or 1)
            return f"{text[:4]}-T{(month - 1) // 3 + 1}"
        if interval == "week":
            try:
                date = datetime.fromisoformat(text)
                return f"{date.isocalendar()[0]}-S{date.isocalendar()[1]:02d}"
            except ValueError:
                return text
        return text
    if isinstance(value, list):
        return [item for item in value]
    return value


def _run_local_query(
    dataset: Dict[str, Any],
    payload: Dict[str, Any],
    dimensions: Sequence[Dict[str, Any]],
    measures: Sequence[Dict[str, Any]],
    formulas: Sequence[Dict[str, Any]],
    filters: Dict[str, Any],
    search: Optional[str],
    limit: int,
    top_n: Optional[int],
    others: bool,
    sort_by: str,
    sort_order: str,
    scope: Optional[Dict[str, Any]],
    notes: List[str],
    started: float,
) -> Dict[str, Any]:
    source_rows: List[Dict[str, Any]] = []
    if dataset["kind"] == "local":
        if dataset.get("requires_session") and not scope:
            return {"error": f"O dataset «{dataset['label']}» exige sessão iniciada.", "rows": [], "meta": {"notes": notes}}
        try:
            source_rows = _local_rows(dataset["id"], (scope or {}).get("user_id"))
        except Exception as exc:
            return {"error": f"Não foi possível carregar os dados: {exc}", "rows": [], "meta": {"notes": notes}}

    filtered: List[Dict[str, Any]] = []
    for row in source_rows:
        if not _local_matches(row, dataset, filters):
            continue
        if isinstance(search, str) and search.strip():
            needle = _normalize(search)
            haystack = " ".join(_normalize(str(value)) for value in row.values() if isinstance(value, (str, int, float)))
            if needle not in haystack:
                continue
        filtered.append(row)

    groups: Dict[Tuple[Any, ...], List[Dict[str, Any]]] = {}
    for row in filtered:
        key = tuple(tuple(_local_bucket(_local_value(row, entry), entry)) if isinstance(_local_bucket(_local_value(row, entry), entry), list)
                    else _local_bucket(_local_value(row, entry), entry) for entry in dimensions)
        groups.setdefault(key, []).append(row)

    rows: List[Dict[str, Any]] = []
    effective_limit = max(limit, 300) if (dimensions and _is_temporal(dimensions[0])) else limit
    truncated = len(groups) > effective_limit
    for key, members in list(groups.items())[: max(effective_limit, 1)]:
        row: Dict[str, Any] = {"_labels": {}}
        for position, entry in enumerate(dimensions):
            row[entry["id"]] = key[position]
            row["_labels"][entry["id"]] = key[position]
        for measure in measures:
            row[measure["id"]] = _local_measure(members, measure)
        row["_key"] = "|".join(str(item) for item in key)
        rows.append(row)

    totals = {measure["id"]: _local_measure(filtered, measure) for measure in measures}
    notes.append(f"{len(filtered)} registo(s) no ficheiro da plataforma; agregação calculada em memória.")
    if truncated:
        notes.append(f"Mostrados {effective_limit} de {len(groups)} grupos (aumente o limite para ver mais).")

    resolution: List[Dict[str, Any]] = [
        {"id": entry["id"], "label": entry.get("label"), "type": entry.get("type"), "field": entry.get("field"),
         "interval": entry.get("interval")}
        for entry in dimensions
    ]
    result = _finalize_rows(rows, dimensions, measures, formulas, resolution, totals, sort_by, sort_order, top_n, others)
    result.update({
        "dataset": _dataset_summary(dataset, scope),
        "totals": totals,
        "meta": {
            "documents": len(filtered),
            "groups": len(groups),
            "elapsed_ms": round((time.time() - started) * 1000, 1),
            "truncated": truncated or result.get("meta", {}).get("truncated", False),
            "notes": notes + list(result.get("meta", {}).get("notes") or []),
            "resolution": resolution,
            "source": "memória",
        },
        "error": None,
    })
    return result


def _local_matches(row: Dict[str, Any], dataset: Dict[str, Any], filters: Dict[str, Any]) -> bool:
    props = {entry["id"]: entry for entry in dataset.get("dimensions") or []}
    for key, value in (filters or {}).items():
        entry = props.get(key)
        if not entry or value in (None, ""):
            continue
        raw = _local_value(row, entry)
        values = raw if isinstance(raw, list) else [raw]
        if isinstance(value, dict):
            numbers = [_to_number(item) for item in values]
            numbers = [number for number in numbers if number is not None]
            low = _to_number(value.get("min", value.get("gte")))
            high = _to_number(value.get("max", value.get("lte")))
            if low is not None and (not numbers or max(numbers) < low):
                return False
            if high is not None and (not numbers or min(numbers) > high):
                return False
        elif isinstance(value, list):
            wanted = {_normalize(item) for item in value}
            if not any(_normalize(item) in wanted for item in values):
                return False
        else:
            if not any(_normalize(item) == _normalize(value) for item in values):
                return False
    return True


# ---------------------------------------------------------------------------
# Fórmulas, ordenação e Top-N
# ---------------------------------------------------------------------------
_ALLOWED_NODES = (
    ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
    ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.Mod, ast.USub, ast.UAdd,
    ast.Call, ast.IfExp, ast.Compare, ast.Lt, ast.LtE, ast.Gt, ast.GtE, ast.Eq, ast.NotEq,
    ast.BoolOp, ast.And, ast.Or,
)
_ALLOWED_FUNCTIONS: Dict[str, Callable[..., Any]] = {
    "abs": abs,
    "round": round,
    "min": min,
    "max": max,
    "sqrt": lambda value: math.sqrt(value) if value is not None and value >= 0 else None,
    "log": lambda value: math.log(value) if value and value > 0 else None,
    "coalesce": lambda *values: next((value for value in values if value is not None), None),
    "ifnull": lambda value, fallback: value if value is not None else fallback,
}

_BRACKET = re.compile(r"\[(@?)([^\[\]]{1,80})\]")


def formula_variables(expression: str, lookup: Callable[[str], Optional[Tuple[str, bool]]]) -> Tuple[str, Dict[str, str]]:
    """Converte `[Nome]` / `[@Nome]` numa expressão Python com variáveis `v0`, `v1`…

    `lookup` recebe o nome escrito pelo utilizador e devolve `(chave, é_total)`.
    """
    mapping: Dict[str, str] = {}
    used: Dict[str, str] = {}

    def replace(match: re.Match) -> str:
        total = bool(match.group(1))
        label = match.group(2).strip()
        found = lookup(label)
        if not found:
            raise ValueError(f"Medida desconhecida na fórmula: [{match.group(0)[1:-1]}]")
        key, is_total = found
        signature = f"{'@' if is_total else ''}{key}"
        if signature in used:
            return used[signature]
        name = f"v{len(mapping)}"
        mapping[name] = f"{'total:' if is_total else ''}{key}"
        used[signature] = name
        return name

    expression = _BRACKET.sub(replace, str(expression or ""))
    if not mapping:
        raise ValueError("A fórmula não usa nenhuma medida. Escreva, por exemplo, [Valor] / [Contagem].")
    return expression, mapping


def evaluate_formula(expression: str, values: Dict[str, Optional[float]]) -> Optional[float]:
    """Avalia uma fórmula já convertida (`v0 / v1`) sobre os valores de uma linha.

    A árvore é validada nó a nó contra uma lista branca (sem atributos, índices,
    lambdas nem acesso a nomes livres) antes de ser executada — as fórmulas vêm do
    cliente e nunca devem poder executar código arbitrário.
    """
    tree = ast.parse(expression, mode="eval")
    forbidden = (ast.Attribute, ast.Subscript, ast.Lambda, ast.ListComp, ast.DictComp,
                 ast.SetComp, ast.GeneratorExp, ast.Import, ast.ImportFrom, ast.NamedExpr)
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            if not isinstance(node.func, ast.Name) or node.func.id not in _ALLOWED_FUNCTIONS:
                raise ValueError("Função não permitida na fórmula.")
            continue
        if isinstance(node, forbidden):
            raise ValueError("A fórmula usa uma construção não permitida.")
        if isinstance(node, (ast.Load, ast.operator, ast.unaryop, ast.cmpop, ast.boolop, ast.expr_context)):
            continue
        if not isinstance(node, _ALLOWED_NODES):
            raise ValueError("A fórmula usa uma construção não permitida.")
    try:
        result = eval(compile(tree, "<formula>", "eval"), {"__builtins__": {}}, {**_ALLOWED_FUNCTIONS, **values})  # noqa: S307 (AST validada acima)
    except ZeroDivisionError:
        return None
    except Exception:
        return None
    if isinstance(result, bool):
        return None
    if isinstance(result, (int, float)):
        if isinstance(result, float) and (math.isnan(result) or math.isinf(result)):
            return None
        return float(result)
    return None


def _build_formula_lookup(measures: Sequence[Dict[str, Any]], formulas: Sequence[Dict[str, Any]]) -> Callable[[str], Optional[Tuple[str, bool]]]:
    index: Dict[str, str] = {}
    for measure in measures:
        index[_normalize(measure["id"])] = measure["id"]
        if measure.get("label"):
            index.setdefault(_normalize(measure["label"]), measure["id"])
    for formula in formulas:
        index.setdefault(_normalize(formula["id"]), formula["id"])
        index.setdefault(_normalize(formula["label"]), formula["id"])

    def lookup(label: str) -> Optional[Tuple[str, bool]]:
        text = label.strip()
        total = text.startswith("@") or _normalize(text).startswith("total ")
        text = text.lstrip("@").strip()
        if _normalize(text).startswith("total "):
            text = text[6:].strip()
        key = index.get(_normalize(text))
        return (key, total) if key else None

    return lookup


def _is_temporal(entry: Dict[str, Any]) -> bool:
    """Dimensões que devem ser mostradas por ordem cronológica, não por valor.

    `Ano` é guardado como número no índice de contratos, pelo que sem esta
    distinção um gráfico por ano sairia ordenado pelo valor (2025, 2024, 2023,
    2026…) em vez de 2012, 2013, …
    """
    return entry.get("type") == "date" or entry.get("id") in ("ano", "ano_mes", "mes", "trimestre", "semana", "dia")


def _finalize_rows(
    rows: List[Dict[str, Any]],
    dimensions: Sequence[Dict[str, Any]],
    measures: Sequence[Dict[str, Any]],
    formulas: Sequence[Dict[str, Any]],
    resolution: Sequence[Dict[str, Any]],
    totals: Dict[str, Optional[float]],
    sort_by: str,
    sort_order: str,
    top_n: Optional[int],
    others: bool,
) -> Dict[str, Any]:
    notes: List[str] = []
    formula_errors: List[str] = []
    lookup = _build_formula_lookup(measures, formulas)

    compiled: List[Tuple[Dict[str, Any], str, Dict[str, str]]] = []
    for formula in formulas:
        try:
            expression, mapping = formula_variables(formula["expression"], lookup)
            compiled.append((formula, expression, mapping))
        except ValueError as exc:
            formula_errors.append(f"{formula['label']}: {exc}")
    valid_formula_ids = {formula["id"] for formula, _expression, _mapping in compiled}

    for row in rows:
        row_values = {measure["id"]: _to_number(row.get(measure["id"])) for measure in measures}
        for formula in formulas:
            if formula["id"] not in valid_formula_ids:
                row[formula["id"]] = None
        for formula, expression, mapping in compiled:
            scope: Dict[str, float] = {}
            for name, reference in mapping.items():
                is_total = reference.startswith("total:")
                key = reference.split(":", 1)[1] if is_total else reference
                base = totals.get(key) if is_total else row_values.get(key)
                if base is None and not is_total:
                    base = next((_to_number(item.get(key)) for item in rows if item.get("_key") == row.get("_key")), None)
                scope[name] = float(base) if isinstance(base, (int, float)) else 0.0
            row[formula["id"]] = evaluate_formula(expression, scope)
    if formula_errors:
        notes.extend(formula_errors)

    sort_key = sort_by or (measures[0]["id"] if measures else "contagem")
    chronological = bool(dimensions) and _is_temporal(dimensions[0]) and not sort_by
    if chronological:
        rows.sort(key=lambda row: str(row.get(dimensions[0]["id"]) or ""))
    else:
        def _sort_value(row: Dict[str, Any]) -> float:
            number = _to_number(row.get(sort_key))
            return number if number is not None else -math.inf

        rows.sort(key=_sort_value, reverse=sort_order != "asc")

    truncated = False
    if top_n and len(rows) > top_n:
        truncated = True
        head, tail = rows[:top_n], rows[top_n:]
        if others:
            merged: Dict[str, Any] = {"_key": "__outros__", "_labels": {entry["id"]: "Outros" for entry in dimensions}}
            for entry in dimensions:
                merged[entry["id"]] = "Outros"
            for measure in measures:
                if MEASURE_KINDS[measure["kind"]].get("additive"):
                    values = [_to_number(item.get(measure["id"])) for item in tail]
                    values = [value for value in values if value is not None]
                    merged[measure["id"]] = float(sum(values)) if values else None
                else:
                    merged[measure["id"]] = None
            if any(not MEASURE_KINDS[measure["kind"]].get("additive") for measure in measures):
                notes.append("No grupo «Outros» só se somam medidas aditivas (contagem e soma) — as restantes ficam vazias, para não inventar valores.")
            for formula, expression, mapping in compiled:
                scope: Dict[str, float] = {}
                for name, reference in mapping.items():
                    is_total = reference.startswith("total:")
                    key = reference.split(":", 1)[1] if is_total else reference
                    base = totals.get(key) if is_total else merged.get(key)
                    scope[name] = float(base) if isinstance(base, (int, float)) else 0.0
                merged[formula["id"]] = evaluate_formula(expression, scope)
            rows = head + [merged]
        else:
            rows = head

    columns = [
        {
            "key": entry["id"],
            "label": entry.get("label") or entry["id"],
            "type": "dimension",
            "data_type": entry.get("type"),
            "unit": entry.get("unit"),
            "interval": entry.get("interval"),
        }
        for entry in dimensions
    ]
    columns.extend({
        "key": measure["id"],
        "label": measure.get("label") or measure["id"],
        "type": "measure",
        "measure_kind": measure["kind"],
        "unit": measure.get("unit"),
        "format": measure.get("format"),
    } for measure in measures)
    columns.extend({
        "key": formula["id"],
        "label": formula["label"],
        "type": "formula",
        "expression": formula["expression"],
        "format": formula.get("format"),
        "unit": formula.get("unit"),
    } for formula in formulas)

    clean_rows = []
    for row in rows:
        clean = {key: value for key, value in row.items() if not key.startswith("_")}
        clean["_key"] = row.get("_key")
        clean["_labels"] = row.get("_labels") or {}
        clean["_series"] = row.get("_series") or []
        clean_rows.append(clean)

    return {
        "dimensions": [{key: value for key, value in entry.items() if key != "binding"} for entry in dimensions],
        "measures": measures,
        "formulas": formulas,
        "columns": columns,
        "rows": clean_rows,
        "meta": {"truncated": truncated, "notes": notes, "resolution": list(resolution)},
    }


# ---------------------------------------------------------------------------
# Registos (drill-through / tabela) e valores distintos
# ---------------------------------------------------------------------------
def run_records(payload: Dict[str, Any], scope: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Linhas individuais de um dataset (drill-through dos visuais e tabela)."""
    dataset_id = str(payload.get("dataset") or "").strip()
    try:
        dataset = get_dataset(dataset_id)
    except KeyError as exc:
        return {"error": str(exc), "items": [], "total": 0}
    size = max(1, min(int(payload.get("size") or 25), 200))
    from_ = max(0, int(payload.get("from") or 0))
    sort = payload.get("sort") or dataset.get("default_sort")
    filters = payload.get("filters") if isinstance(payload.get("filters"), dict) else {}

    if dataset.get("requires_session") and not scope:
        return {"error": f"O dataset «{dataset['label']}» exige sessão iniciada.", "items": [], "total": 0}

    if dataset["kind"] == "local":
        try:
            rows = _local_rows(dataset["id"], (scope or {}).get("user_id"))
        except Exception as exc:
            return {"error": f"Não foi possível carregar os dados: {exc}", "items": [], "total": 0}
        filtered = [row for row in rows if _local_matches(row, dataset, filters)]
        return {
            "dataset": _dataset_summary(dataset, scope),
            "total": len(filtered),
            "items": filtered[from_: from_ + size],
            "columns": [{"key": entry["id"], "label": entry.get("label")} for entry in dataset.get("dimensions") or []],
            "error": None,
        }

    try:
        result = ontology.query_objects(
            dataset_id,
            search=payload.get("search"),
            filters=filters,
            size=size,
            from_=from_,
            sort=sort,
            scope=scope,
        )
    except PermissionError as exc:
        return {"error": str(exc), "items": [], "total": 0}
    except KeyError as exc:
        return {"error": str(exc), "items": [], "total": 0}
    items = result.get("items") or []
    columns = list(items[0].keys()) if items else [entry["id"] for entry in dataset.get("dimensions") or []]
    return {
        "dataset": _dataset_summary(dataset, scope),
        "total": result.get("total", 0),
        "items": items,
        "columns": [{"key": key, "label": key} for key in columns if key not in ("_score",)],
        "exact": result.get("exact", True),
        "notes": result.get("notes") or [],
        "error": result.get("error"),
    }


def distinct_values(payload: Dict[str, Any], scope: Optional[Dict[str, Any]] = None, es: Any = None) -> Dict[str, Any]:
    """Valores distintos de uma dimensão (para os seletores de filtros)."""
    dataset_id = str(payload.get("dataset") or "").strip()
    field_id = str(payload.get("field") or "").strip()
    needle = _normalize(payload.get("q") or "")
    limit = max(1, min(int(payload.get("limit") or 50), 200))
    try:
        dataset = get_dataset(dataset_id)
    except KeyError as exc:
        return {"error": str(exc), "items": []}
    entry = next((item for item in dataset.get("dimensions") or [] if item["id"] == field_id), None)
    if not entry:
        return {"error": f"Dimensão desconhecida: {field_id}", "items": []}
    if dataset.get("requires_session") and not scope:
        return {"error": f"O dataset «{dataset['label']}» exige sessão iniciada.", "items": []}

    if dataset["kind"] == "local":
        try:
            rows = _local_rows(dataset["id"], (scope or {}).get("user_id"))
        except Exception as exc:
            return {"error": str(exc), "items": []}
        values: List[Any] = []
        for row in rows:
            raw = _local_value(row, entry)
            for item in (raw if isinstance(raw, list) else [raw]):
                if item not in (None, "") and item not in values:
                    values.append(item)
        values = [item for item in values if not needle or needle in _normalize(item)]
        return {
            "dataset": dataset_id,
            "field": field_id,
            "total": len(values),
            "items": values[:limit],
            "truncated": len(values) > limit,
            "error": None,
        }

    client = es or get_es_client()
    if not client:
        return {"error": "Elasticsearch indisponível.", "items": []}
    index = _index_of_dataset(dataset)
    properties = _index_properties(client, index or "")
    runtime: Dict[str, Any] = {}
    target = _agg_target(client, index, entry["field"], properties)
    if target.get("runtime"):
        runtime.update(target["runtime"])
    try:
        query, _ = _base_query(dataset, payload.get("filters") or {}, payload.get("search"))
        scope_clauses = _scope_for(scope, dataset)
    except PermissionError as exc:
        return {"error": str(exc), "items": []}
    if scope_clauses:
        query.setdefault("bool", {}).setdefault("filter", []).extend(scope_clauses)

    node: Dict[str, Any] = {"terms": {"field": target["field"], "size": max(limit * 4, 100)}}
    wrapper: Dict[str, Any] = {"valores": node}
    if entry.get("nested"):
        wrapper = {"aninhado": {"nested": {"path": entry["nested"]}, "aggs": wrapper}}
    body: Dict[str, Any] = {"size": 0, "query": query, "aggs": wrapper}
    if runtime:
        body["runtime_mappings"] = runtime
    try:
        response = client.search(index=index, body=body, request_timeout=AGG_TIMEOUT_SECONDS)
    except Exception as exc:
        return {"error": f"A consulta falhou: {exc}", "items": []}
    buckets = (((response.get("aggregations") or {}).get("aninhado") or {}).get("valores")
               or (response.get("aggregations") or {}).get("valores") or {}).get("buckets") or []
    items = [bucket.get("key") for bucket in buckets if bucket.get("key") not in (None, "")]
    if needle:
        items = [item for item in items if needle in _normalize(item)]
    return {
        "dataset": dataset_id,
        "field": field_id,
        "total": len(items),
        "items": items[:limit],
        "truncated": len(items) > limit,
        "error": None,
    }


# ---------------------------------------------------------------------------
# Exportação
# ---------------------------------------------------------------------------
def _table_from_result(result: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], List[Dict[str, Any]]]:
    columns = result.get("columns") or []
    headers = [column.get("label") or column["key"] for column in columns]
    rows: List[List[Any]] = []
    for row in result.get("rows") or []:
        rows.append([row.get(column["key"]) for column in columns])
    return headers, rows, columns


def export_csv(result: Dict[str, Any], dataset_label: str = "dados") -> Tuple[str, bytes, str]:
    """Exporta o resultado para CSV (separador `;`, pronto para Excel pt-PT)."""
    headers, rows, columns = _table_from_result(result)
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    meta_lines = [
        ["# Visualizador — IQ OS"],
        ["# Dataset", dataset_label],
        ["# Gerado em", _now_iso()],
    ]
    if result.get("meta", {}).get("notes"):
        meta_lines.append(["# Notas", " | ".join(result["meta"]["notes"])])
    for line in meta_lines:
        writer.writerow(line)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_csv_value(value, column) for value, column in zip(row, columns)])
    filename = f"visualizador_{_slugify(dataset_label)}_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    content = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    return filename, content, "text/csv; charset=utf-8"


def _csv_value(value: Any, column: Dict[str, Any]) -> Any:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.6f}".rstrip("0").rstrip(".").replace(".", ",")
    return value


def export_xlsx(result: Dict[str, Any], dataset_label: str = "dados") -> Tuple[str, bytes, str]:
    """Exporta o resultado para Excel, com folha de dados e folha de definição."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:  # openpyxl em falta
        raise RuntimeError("A exportação para Excel requer a biblioteca openpyxl.") from exc

    headers, rows, columns = _table_from_result(result)
    book = Workbook()
    sheet = book.active
    sheet.title = "Dados"
    header_fill = PatternFill("solid", fgColor="10A37F")
    for position, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=position, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            column = columns[column_index - 1]
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.value = value
                if column.get("format") == "currency" or column.get("unit") == "€":
                    cell.number_format = '#,##0.00 €'
                elif isinstance(value, float):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.value = "" if value is None else value
    sheet.freeze_panes = "A2"
    for column_index, header in enumerate(headers, start=1):
        width = max(12, min(48, max([len(str(header))] + [len(str(row[column_index - 1] or "")) for row in rows[:200]] if rows else [len(str(header))]) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    info = book.create_sheet("Definição")
    info_rows = [
        ("Dataset", dataset_label),
        ("Gerado em", _now_iso()),
        ("Registos analisados", result.get("meta", {}).get("documents")),
        ("Grupos", len(rows)),
        ("Truncado", "sim" if result.get("meta", {}).get("truncated") else "não"),
    ]
    for measure in result.get("measures") or []:
        info_rows.append(("Medida", f"{measure.get('label')} ({measure.get('kind')} de {measure.get('field') or 'registos'})"))
    for formula in result.get("formulas") or []:
        info_rows.append(("Fórmula", f"{formula.get('label')} = {formula.get('expression')}"))
    for note in (result.get("meta", {}).get("notes") or []):
        info_rows.append(("Nota", note))
    for row_index, (label, value) in enumerate(info_rows, start=1):
        info.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        info.cell(row=row_index, column=2, value=value)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    buffer = io.BytesIO()
    book.save(buffer)
    filename = f"visualizador_{_slugify(dataset_label)}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return filename, buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _slugify(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()[:48] or "dados"
