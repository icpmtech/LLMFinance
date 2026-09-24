"""Relatórios do módulo **Contribuintes** — PDF, Excel e CSV com a marca do IQ OS.

O que produz
------------
- **Ficha de um contribuinte** (`contribuinte_report`): identificação, tipo,
  país, localização, totais de atividade (contratos PT/ES, publicações
  societárias, CIRE, marcas, firmas, cargos, CRM) e o detalhe por fonte
  (`src_*`), incluindo as designações alternativas.
- **Lista de contribuintes** (`list_report`): indicadores agregados do conjunto
  filtrado (contagem por tipo, valor contratual, localização) e a tabela dos
  contribuintes.

Como
----
Uma única estrutura de **secções** (`{title, rows: [[rótulo, valor]]}`) alimenta
os três renderizadores, para que PDF, Excel e CSV digam sempre o mesmo:

- **PDF** — ReportLab (Platypus) com o logótipo do IQ OS no cabeçalho;
- **Excel** — openpyxl, com o logótipo embutido e colunas dimensionadas;
- **CSV** — texto com o cabeçalho da marca, separador `;` e BOM UTF-8 (abre
  corretamente no Excel português).

O CSV não é um formato binário: não pode conter a imagem, pelo que leva a marca
em texto. O logótipo vive em `chat-ui/public/` (o mesmo que a aplicação usa).
"""
from __future__ import annotations

import logging
import re
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from api.elasticsearch_client import ROOT

logger = logging.getLogger(__name__)

BRAND = "IQ OS"
BRAND_TAGLINE = "Plataforma de Inteligência Financeira"

#: Formatos suportados (id → etiqueta e tipo MIME).
FORMATS: Dict[str, Dict[str, str]] = {
    "pdf": {"label": "PDF", "media_type": "application/pdf", "extension": "pdf"},
    "xlsx": {
        "label": "Excel",
        "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "extension": "xlsx",
    },
    "csv": {"label": "CSV", "media_type": "text/csv; charset=utf-8", "extension": "csv"},
}

_CURRENCY_KEYS = ("value", "valor", "valor_contratual", "entities_value", "contracts_value")

# ---------------------------------------------------------------------- marca
def logo_path() -> Optional[Path]:
    """Caminho do logótipo do IQ OS (o mesmo ícone que a aplicação usa)."""
    candidates = [
        ROOT / "chat-ui" / "public" / "icon-192.png",
        ROOT / "chat-ui" / "public" / "icon-512.png",
        ROOT / "chat-ui" / "public" / "apple-touch-icon.png",
        ROOT / "chat-ui" / "public" / "favicon-32.png",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def slugify(value: str, *, fallback: str = "contribuintes") -> str:
    """Nome de ficheiro seguro (sem acentos, espaços ou símbolos).

    O NFKD separa a letra do acento (`Ê` → `E` + marca combinante), pelo que as
    marcas são removidas explicitamente — sem isso, «FARMACÊUTICA» saía como
    «farmace-utica» (a marca combinante virava separador).
    """
    decomposed = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in decomposed if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").lower()
    return text[:60] or fallback


def _pretty_label(key: str) -> str:
    """Chave de um detalhe → rótulo legível.

    `natureza_juridica` → «Natureza juridica»; `localExecucao` → «Local execucao».
    """
    text = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", str(key or "")).replace("_", " ").strip()
    if not text:
        return str(key or "")
    return text[:1].upper() + text[1:].lower()


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d")


# ----------------------------------------------------------------- formatação
def _number(value: Any) -> str:
    """Número no formato português (o que o leitor espera no relatório)."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if number.is_integer():
        return f"{int(number):,}".replace(",", " ")
    return f"{number:,.2f}".replace(",", " ").replace(".", ",")


def _money(value: Any) -> str:
    try:
        return f"{_number(float(value))} €"
    except (TypeError, ValueError):
        return str(value)


def _date(value: Any) -> str:
    if not value:
        return "—"
    text = str(value)
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        year, month, day = match.groups()
        return f"{day}/{month}/{year}"
    return text


def _present(value: Any) -> str:
    """Valor pronto para o relatório (nunca vazio nem dicionários crus)."""
    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return "sim" if value else "não"
    if isinstance(value, (int, float)):
        return _number(value)
    if isinstance(value, (list, tuple, set)):
        items = [str(item) for item in value if item not in (None, "")]
        return ", ".join(items) if items else "—"
    if isinstance(value, dict):
        pairs = [f"{key}: {_present(item)}" for key, item in value.items() if item not in (None, "")]
        return "; ".join(pairs) if pairs else "—"
    return str(value)


def _is_currency_key(key: str) -> bool:
    return any(token in key.lower() for token in _CURRENCY_KEYS)


# -------------------------------------------------------------------- secções
def _source_blocks(doc: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    """Blocos de evidência por fonte (`src_*`), com uma etiqueta legível."""
    blocks: List[Tuple[str, Dict[str, Any]]] = []
    for key in sorted(doc.keys()):
        if not key.startswith("src_"):
            continue
        block = doc.get(key)
        if not isinstance(block, dict):
            continue
        source_id = key[4:]
        label = str(block.get("label") or "").strip()
        if not label:
            label = source_id.replace("_", " ").capitalize()
        blocks.append((label, block))
    return blocks


def _block_rows(block: Dict[str, Any]) -> List[List[str]]:
    """Linhas de evidência de uma fonte (contagens, valores, datas, detalhes)."""
    rows: List[List[str]] = []
    if block.get("count") is not None:
        rows.append(["Registos", _number(block.get("count"))])
    if block.get("value"):
        rows.append(["Valor", _money(block.get("value"))])
    if block.get("first") or block.get("last"):
        rows.append(["Período", f"{_date(block.get('first'))} → {_date(block.get('last'))}"])
    if block.get("country"):
        rows.append(["País", _present(block.get("country"))])
    if block.get("roles"):
        rows.append(["Papéis", _present(block.get("roles"))])
    names = [str(name) for name in (block.get("names") or []) if name]
    if names:
        rows.append(["Designações", "; ".join(names[:6])])
    parts = block.get("parts") or {}
    if isinstance(parts, dict):
        for name, part in parts.items():
            if not isinstance(part, dict):
                continue
            count = part.get("count")
            value = part.get("value")
            detail = f"{_number(count)}" if count is not None else "—"
            if value:
                detail += f" · {_money(value)}"
            if part.get("last"):
                detail += f" · última {_date(part.get('last'))}"
            rows.append([str(name).replace("_", " "), detail])
    detail = block.get("detail") or {}
    if isinstance(detail, dict):
        for name, value in detail.items():
            # Listas (ex.: `localExecucao`) são mostradas juntas; dicionários
            # aninhados não cabem numa linha da tabela.
            if value in (None, "") or isinstance(value, dict):
                continue
            rows.append([_pretty_label(str(name)), _money(value) if _is_currency_key(str(name)) else _present(value)])
    return rows


def contribuinte_sections(doc: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Secções da ficha de um contribuinte (a matéria-prima dos três formatos)."""
    sections: List[Dict[str, Any]] = []

    identification = [
        ["NIF / NIPC", _present(doc.get("nif"))],
        ["Designação", _present(doc.get("name"))],
        ["Tipo de contribuinte", _present(doc.get("type_label") or doc.get("type"))],
        ["País", _present(doc.get("country"))],
        ["Natureza", "Pessoa coletiva" if doc.get("is_company") else "Pessoa singular"],
        ["Dígito de controlo", "válido" if doc.get("nif_valid") else "a confirmar"],
        ["Fontes", _present(doc.get("source_labels") or doc.get("sources"))],
        ["Papéis", _present(doc.get("roles"))],
        ["Primeiro registo", _date(doc.get("first_seen"))],
        ["Último registo", _date(doc.get("last_seen"))],
        ["Registos agregados", _number(doc.get("records_total") or 0)],
    ]
    sections.append({"title": "Identificação", "rows": identification})

    location = doc.get("location") or {}
    if isinstance(location, dict) and any(location.values()):
        labels = {
            "pais": "País",
            "distrito": "Distrito",
            "concelho": "Concelho",
            "freguesia": "Freguesia",
            "codigo_postal": "Código postal",
        }
        rows = [[labels.get(key, key), _present(value)] for key, value in location.items() if value]
        if rows:
            sections.append({"title": "Localização", "rows": rows})

    activity: List[List[str]] = []
    counters = [
        ("Contratos públicos (PT)", doc.get("contracts_count")),
        ("· como adjudicante", doc.get("contracts_as_adjudicante")),
        ("· como adjudicatário", doc.get("contracts_as_adjudicatario")),
        ("Contratos públicos (ES)", doc.get("contratos_es_count")),
        ("Contratos no cadastro", doc.get("entities_contracts_count")),
        ("Publicações societárias", doc.get("societario_count")),
        ("Processos CIRE", doc.get("cire_count")),
        ("Marcas (INPI)", doc.get("trademarks_count")),
        ("Firmas (RNPC)", doc.get("firmas_count")),
        ("Cargos (PessoasIQ)", doc.get("people_roles_count")),
        ("Empresas ligadas", doc.get("people_companies_count")),
    ]
    for label, value in counters:
        if value:
            activity.append([label, _number(value)])
    values = [
        ("Valor contratual (PT)", doc.get("contracts_value")),
        ("Valor contratual (ES)", doc.get("contratos_es_value")),
        ("Valor no cadastro", doc.get("entities_value")),
    ]
    for label, value in values:
        if value:
            activity.append([label, _money(value)])
    dates = [
        ("Contratos PT — última", doc.get("contracts_last_date")),
        ("Contratos ES — última", doc.get("contratos_es_last_date")),
        ("Publicações — última", doc.get("societario_last_date")),
        ("CIRE — última", doc.get("cire_last_date")),
    ]
    for label, value in dates:
        if value:
            activity.append([label, _date(value)])
    if doc.get("cire_roles"):
        activity.append(["Papéis no CIRE", _present(doc.get("cire_roles"))])
    if doc.get("crm_account"):
        activity.append(["Conta de CRM", "sim"])
    if activity:
        sections.append({"title": "Atividade agregada", "rows": activity})

    names = [str(name) for name in (doc.get("names") or []) if name]
    if len(names) > 1:
        sections.append({"title": "Designações conhecidas", "rows": [[f"Designação {index + 1}", name] for index, name in enumerate(names)]})

    for label, block in _source_blocks(doc):
        rows = _block_rows(block)
        if rows:
            sections.append({"title": f"Fonte: {label}", "rows": rows})

    if doc.get("synced_at"):
        sections.append(
            {
                "title": "Controlo",
                "rows": [
                    ["Última sincronização", _date(doc.get("synced_at"))],
                    ["Passagem (run_id)", _present(doc.get("run_id"))],
                ],
            }
        )
    return sections


#: Colunas do relatório de lista (o que sai no PDF, Excel e CSV).
LIST_COLUMNS: Sequence[Tuple[str, str]] = (
    ("nif", "NIF"),
    ("name", "Designação"),
    ("type_label", "Tipo"),
    ("country", "País"),
    ("local", "Localização"),
    ("sources", "Fontes"),
    ("roles", "Papéis"),
    ("contracts_count", "Contratos"),
    ("contracts_value", "Valor contratual"),
    ("last_seen", "Último registo"),
)


def list_row(item: Dict[str, Any]) -> Dict[str, str]:
    """Achata um contribuinte nas colunas do relatório de lista."""
    location = item.get("location") or {}
    local = " · ".join(
        str(location.get(key)) for key in ("concelho", "distrito") if isinstance(location, dict) and location.get(key)
    )
    return {
        "nif": str(item.get("nif") or ""),
        "name": str(item.get("name") or ""),
        "type_label": str(item.get("type_label") or item.get("type") or ""),
        "country": str(item.get("country") or ""),
        "local": local,
        "sources": ", ".join(str(value) for value in (item.get("source_labels") or item.get("sources") or [])),
        "roles": ", ".join(str(value) for value in (item.get("roles") or [])),
        "contracts_count": _number(item.get("contracts_count") or 0),
        "contracts_value": _money(item.get("contracts_value") or 0),
        "last_seen": _date(item.get("last_seen")),
    }


def list_sections(items: List[Dict[str, Any]], *, total: int, filters_label: str = "") -> List[Dict[str, Any]]:
    """Secções do relatório de lista (indicadores + tabela de contribuintes)."""
    by_type: Dict[str, int] = {}
    for item in items:
        key = str(item.get("type_label") or item.get("type") or "desconhecido")
        by_type[key] = by_type.get(key, 0) + 1
    value_total = 0.0
    for item in items:
        try:
            value_total += float(item.get("contracts_value") or 0)
        except (TypeError, ValueError):
            continue
    with_contracts = sum(1 for item in items if (item.get("contracts_count") or 0) > 0)
    with_location = sum(1 for item in items if isinstance(item.get("location"), dict) and any((item.get("location") or {}).values()))

    summary: List[List[str]] = [
        ["Contribuintes no resultado", _number(total)],
        ["Contribuintes exportados", _number(len(items))],
        ["Com contratos públicos", _number(with_contracts)],
        ["Com localização conhecida", _number(with_location)],
        ["Valor contratual (exportado)", _money(value_total)],
    ]
    if filters_label:
        summary.append(["Filtros aplicados", filters_label])
    sections: List[Dict[str, Any]] = [{"title": "Resumo", "rows": summary}]

    if by_type:
        ordered = sorted(by_type.items(), key=lambda pair: pair[1], reverse=True)
        sections.append({"title": "Distribuição por tipo", "rows": [[key, _number(count)] for key, count in ordered]})

    sections.append(
        {
            "title": "Contribuintes",
            "columns": [label for _, label in LIST_COLUMNS],
            "rows": [[list_row(item)[key] for key, _ in LIST_COLUMNS] for item in items],
        }
    )
    return sections


# --------------------------------------------------------------------------- CSV
def render_csv(*, title: str, subtitle: str, sections: List[Dict[str, Any]]) -> bytes:
    """CSV com o cabeçalho da marca (separador `;`, pronto para o Excel PT)."""
    separator = ";"

    def escape(value: Any) -> str:
        text = "" if value is None else str(value)
        if any(char in text for char in '";\n'):
            return '"' + text.replace('"', '""') + '"'
        return text

    lines: List[str] = [
        f"# {BRAND} — {BRAND_TAGLINE}",
        f"# {title}",
        f"# {subtitle}",
        f"# Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        "",
    ]
    for section in sections:
        lines.append(f"## {section['title']}")
        columns = section.get("columns")
        if columns:
            lines.append(separator.join(escape(column) for column in columns))
            for row in section.get("rows") or []:
                lines.append(separator.join(escape(cell) for cell in row))
        else:
            lines.append(f"Campo{separator}Valor")
            for row in section.get("rows") or []:
                lines.append(separator.join(escape(cell) for cell in row))
        lines.append("")
    return ("\ufeff" + "\n".join(lines)).encode("utf-8")


# ------------------------------------------------------------------------- Excel
def render_xlsx(*, title: str, subtitle: str, sections: List[Dict[str, Any]]) -> bytes:
    """Livro Excel com o logótipo do IQ OS e uma folha por bloco de dados."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependência declarada
        raise RuntimeError("openpyxl não instalado") from exc

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="0F172A")
    brand_font = Font(bold=True, size=11, color="0F766E")
    muted_font = Font(italic=True, size=9, color="64748B")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    workbook = Workbook()

    def write_sheet(sheet: Any, blocks: List[Dict[str, Any]], *, start_row: int, with_table: bool) -> int:
        row = start_row
        for block in blocks:
            sheet.cell(row=row, column=1, value=block["title"]).font = Font(bold=True, size=11, color="0F172A")
            row += 1
            columns = block.get("columns")
            if columns:
                for index, column in enumerate(columns, start=1):
                    cell = sheet.cell(row=row, column=index, value=column)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")
                row += 1
                for data_row in block.get("rows") or []:
                    for index, value in enumerate(data_row, start=1):
                        cell = sheet.cell(row=row, column=index, value=value)
                        cell.border = border
                        cell.alignment = Alignment(vertical="top", wrap_text=index == 2)
                    row += 1
            else:
                for label, value in block.get("rows") or []:
                    label_cell = sheet.cell(row=row, column=1, value=label)
                    label_cell.font = Font(bold=True, size=10, color="334155")
                    label_cell.border = border
                    value_cell = sheet.cell(row=row, column=2, value=value)
                    value_cell.border = border
                    value_cell.alignment = Alignment(vertical="top", wrap_text=True)
                    row += 1
            row += 1
        if with_table:
            max_columns = max(
                [len(block.get("columns") or []) or 2 for block in blocks] or [2]
            )
            for index in range(1, max_columns + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 34 if index == 2 else 20
        return row

    main = workbook.active
    main.title = "Relatório"
    main["A1"] = BRAND
    main["A1"].font = title_font
    main["A2"] = BRAND_TAGLINE
    main["A2"].font = brand_font
    main["A3"] = title
    main["A3"].font = Font(bold=True, size=12)
    main["A4"] = subtitle
    main["A4"].font = muted_font
    write_sheet(main, sections, start_row=6, with_table=True)

    logo = logo_path()
    if logo:
        try:
            from openpyxl.drawing.image import Image as XlImage

            picture = XlImage(str(logo))
            picture.width = 48
            picture.height = 48
            main.add_image(picture, "D1")
        except Exception as exc:  # pragma: no cover - depende do Pillow
            logger.debug("Logótipo não embutido no Excel: %s", exc)

    # As listas longas ganham uma folha própria, com a primeira linha fixa.
    for block in sections:
        if not block.get("columns") or len(block.get("rows") or []) < 2:
            continue
        sheet = workbook.create_sheet(str(block["title"])[:28])
        write_sheet(sheet, [block], start_row=1, with_table=True)
        sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF
def render_pdf(*, title: str, subtitle: str, sections: List[Dict[str, Any]]) -> bytes:
    """PDF com o logótipo do IQ OS, secções em tabela e rodapé com a marca."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Image as PdfImage,
            KeepTogether,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover - dependência declarada
        raise RuntimeError("reportlab não instalado") from exc

    brand_color = colors.HexColor("#0F766E")
    dark = colors.HexColor("#0F172A")
    muted = colors.HexColor("#64748B")
    zebra = colors.HexColor("#F1F5F9")

    base = getSampleStyleSheet()
    title_style = ParagraphStyle("iq-title", parent=base["Title"], fontSize=17, leading=20, textColor=dark, alignment=0)
    brand_style = ParagraphStyle("iq-brand", parent=base["Normal"], fontSize=10, leading=12, textColor=brand_color, fontName="Helvetica-Bold")
    sub_style = ParagraphStyle("iq-sub", parent=base["Normal"], fontSize=9, leading=12, textColor=muted)
    section_style = ParagraphStyle("iq-section", parent=base["Normal"], fontSize=11, leading=14, textColor=dark, fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=4)
    cell_style = ParagraphStyle("iq-cell", parent=base["Normal"], fontSize=8, leading=10, textColor=dark)
    cell_muted = ParagraphStyle("iq-cell-muted", parent=cell_style, textColor=muted, fontName="Helvetica-Bold")

    def paragraph(value: Any, style: Any = cell_style) -> Paragraph:
        text = str(value if value is not None else "")
        escape = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(escape or "—", style)

    story: List[Any] = []
    logo = logo_path()
    if logo:
        try:
            picture = PdfImage(str(logo), width=13 * mm, height=13 * mm)
            header = Table(
                [[picture, [Paragraph(title, title_style), Paragraph(subtitle, sub_style)]]],
                colWidths=[16 * mm, None],
            )
            header.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (0, 0), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
            story.append(header)
        except Exception as exc:  # pragma: no cover
            logger.debug("Logótipo não embutido no PDF: %s", exc)
            story.append(Paragraph(title, title_style))
    else:
        story.append(Paragraph(title, title_style))
    story.append(Paragraph(f"{BRAND} · {BRAND_TAGLINE}", brand_style))
    story.append(Spacer(1, 8))

    table_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), brand_color),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8.5),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, zebra]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
    )

    max_rows_per_section = 400
    for section in sections:
        columns = section.get("columns")
        rows = section.get("rows") or []
        if not rows:
            continue
        story.append(Paragraph(str(section["title"]), section_style))
        if columns:
            data = [[paragraph(column, ParagraphStyle("h", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")) for column in columns]]
            for row in rows[:max_rows_per_section]:
                data.append([paragraph(cell) for cell in row])
            widths = [None] * len(columns)
            table = Table(data, colWidths=widths, repeatRows=1)
        else:
            data = [[paragraph("Campo", cell_muted), paragraph("Valor", cell_muted)]]
            for label, value in rows[:max_rows_per_section]:
                data.append([paragraph(label, cell_muted), paragraph(value)])
            table = Table(data, colWidths=[55 * mm, None], repeatRows=1)
        table.setStyle(table_style)
        story.append(KeepTogether(table) if len(data) <= 14 else table)
        story.append(Spacer(1, 8))

    story.append(Paragraph(f"Gerado pelo {BRAND} em {datetime.now().strftime('%d/%m/%Y às %H:%M')}.", sub_style))

    buffer = BytesIO()

    def decorate(canvas: Any, doc: Any) -> None:
        """Rodapé com a marca e a numeração (todas as páginas)."""
        canvas.saveState()
        width, _ = doc.pagesize
        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        canvas.setLineWidth(0.5)
        canvas.line(18, 30, width - 18, 30)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(18, 20, f"{BRAND} — {BRAND_TAGLINE}")
        canvas.drawRightString(width - 18, 20, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4) if any(section.get("columns") for section in sections) else A4,
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=38,
        title=f"{title} · {BRAND}",
        author=BRAND,
    )
    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


# -------------------------------------------------------------------- fachada
def _render(format: str, *, title: str, subtitle: str, sections: List[Dict[str, Any]]) -> Tuple[str, bytes, str]:
    fmt = (format or "pdf").strip().lower()
    if fmt not in FORMATS:
        raise ValueError(f"Formato não suportado: {format!r} (use pdf, xlsx ou csv).")
    if fmt == "csv":
        content = render_csv(title=title, subtitle=subtitle, sections=sections)
    elif fmt == "xlsx":
        content = render_xlsx(title=title, subtitle=subtitle, sections=sections)
    else:
        content = render_pdf(title=title, subtitle=subtitle, sections=sections)
    return fmt, content, FORMATS[fmt]["media_type"]


def contribuinte_report(doc: Dict[str, Any], format: str = "pdf") -> Tuple[str, bytes, str]:
    """Relatório de **um** contribuinte. Devolve `(nome do ficheiro, bytes, mime)`."""
    nif = str(doc.get("nif") or "").strip()
    name = str(doc.get("name") or "").strip()
    title = f"Ficha de contribuinte · {name or nif}"
    subtitle = f"NIF {nif}" + (f" · {name}" if name and name != nif else "")
    fmt, content, media = _render(format, title=title, subtitle=subtitle, sections=contribuinte_sections(doc))
    filename = f"iq-os-contribuinte-{slugify(nif or name)}_{_stamp()}.{FORMATS[fmt]['extension']}"
    return filename, content, media


def list_report(
    items: Iterable[Dict[str, Any]],
    *,
    total: int,
    format: str = "pdf",
    filters_label: str = "",
) -> Tuple[str, bytes, str]:
    """Relatório da **lista** de contribuintes. Devolve `(nome do ficheiro, bytes, mime)`."""
    rows = list(items)
    title = "Contribuintes · relatório de resultados"
    subtitle = f"{_number(len(rows))} de {_number(total)} contribuintes"
    if filters_label:
        subtitle += f" · {filters_label}"
    fmt, content, media = _render(
        format,
        title=title,
        subtitle=subtitle,
        sections=list_sections(rows, total=total, filters_label=filters_label),
    )
    filename = f"iq-os-contribuintes_{_stamp()}.{FORMATS[fmt]['extension']}"
    return filename, content, media


def available() -> Dict[str, Any]:
    """Dependências e marca disponíveis (para `/contribuintes/meta`)."""
    pdf = xlsx = False
    try:
        import reportlab  # noqa: F401

        pdf = True
    except Exception:
        pass
    try:
        import openpyxl  # noqa: F401

        xlsx = True
    except Exception:
        pass
    logo = logo_path()
    return {
        "formats": [
            {
                "id": key,
                "label": entry["label"],
                "media_type": entry["media_type"],
                "available": pdf if key == "pdf" else (xlsx if key == "xlsx" else True),
                "embeds_logo": key in ("pdf", "xlsx"),
            }
            for key, entry in FORMATS.items()
        ],
        "brand": BRAND,
        "tagline": BRAND_TAGLINE,
        "logo": str(logo) if logo else "",
        "logo_available": bool(logo),
    }
