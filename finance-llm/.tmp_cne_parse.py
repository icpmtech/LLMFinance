"""Protótipo do parser do Mapa Oficial da CNE — validação contra os mandatos do MAI (temporário)."""
import json
import re
import unicodedata

import openpyxl

CAMINHO = ".tmp_cne/2025al-mapa-oficial_retificado/mapa_3_eleitos_retificado.xlsx"

livro = openpyxl.load_workbook(CAMINHO, read_only=True, data_only=True)
folha = livro[livro.sheetnames[0]]
linhas = list(folha.iter_rows(min_row=1, values_only=True))
print("linhas:", len(linhas))


def vazio(valor):
    return valor is None or str(valor).strip() == ""


# Agrupa em blocos: uma linha com CÓD em col.0 é cabeçalho de bloco.
blocos = []
atual = None
for linha in linhas:
    colunas = list(linha)
    if len(colunas) >= 4 and not vazio(colunas[0]) and not vazio(colunas[3]):
        if atual:
            blocos.append(atual)
        atual = {"cod": str(colunas[0]).strip(), "concelho": str(colunas[1]).strip(), "freguesia": str(colunas[2]).strip() if not vazio(colunas[2]) else None, "org": str(colunas[3]).strip(), "linhas": []}
    elif atual is not None:
        atual["linhas"].append([None if vazio(c) else str(c).strip() for c in colunas])
if atual:
    blocos.append(atual)
print("blocos:", len(blocos))
from collections import Counter

print("por órgão:", Counter(b["org"] for b in blocos))

cm = [b for b in blocos if b["org"] == "CM"]
print("blocos CM:", len(cm))

# Analisa um bloco CM: linha 0 = siglas; seguintes = nomes por coluna
def analisar(bloco):
    if not bloco["linhas"]:
        return None
    siglas = bloco["linhas"][0]
    listas = [{"sigla": sigla, "nomes": []} for sigla in siglas if sigla]
    for linha in bloco["linhas"][1:]:
        for indice, nome in enumerate(linha):
            if indice < len(listas) and nome:
                listas[indice]["nomes"].append(nome)
    return listas


exemplo = analisar(cm[0])
print()
print("exemplo", cm[0]["concelho"], cm[0]["cod"])
for lista in exemplo:
    print("  ", lista["sigla"], "->", len(lista["nomes"]), ":", lista["nomes"][:2])

# Integridade: a soma dos nomes tem de dar os mandatos da câmara
mandatos = {}
import httpx

dados_mai = json.load(open("data/remuneracoes/mai_eleitos.json", encoding="utf-8"))
for registo in dados_mai["concelhos"].values():
    mandatos[registo["nome"]] = sum(p.get("mandatos") or 0 for p in (registo.get("camara") or []))


def chave(texto):
    texto = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", texto).strip()


ok = 0
divergentes = []
for bloco in cm:
    listas = analisar(bloco)
    if not listas:
        continue
    total = sum(len(lista["nomes"]) for lista in listas)
    esperado = mandatos.get(bloco["concelho"]) or mandatos.get(bloco["concelho"].title())
    if esperado is None:
        # procura ignorando maiúsculas
        for nome, valor in mandatos.items():
            if chave(nome) == chave(bloco["concelho"]):
                esperado = valor
                break
    if esperado is None:
        divergentes.append((bloco["concelho"], "sem correspondência", total))
    elif esperado == total:
        ok += 1
    else:
        divergentes.append((bloco["concelho"], esperado, total))

print()
print("blocos CM com soma de nomes == mandatos:", ok, "de", len(cm))
print("divergentes:", len(divergentes), divergentes[:10])

# nomes truncados?
tamanhos = [len(nome) for b in cm for linha in b["linhas"][1:] for nome in linha if nome]
print("nome mais longo:", max(tamanhos), "| nº com >=40 caracteres:", sum(1 for t in tamanhos if t >= 40))
print("exemplo longo:", [nome for b in cm for linha in b["linhas"][1:] for nome in linha if nome and len(nome) >= 40][:3])
