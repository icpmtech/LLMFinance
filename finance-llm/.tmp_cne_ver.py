"""Inspecciona o mapa_3_eleitos (nomes dos eleitos) do Mapa Oficial da CNE (temporário)."""
import openpyxl

CAMINHO = ".tmp_cne/2025al-mapa-oficial_retificado/mapa_3_eleitos_retificado.xlsx"
livro = openpyxl.load_workbook(CAMINHO, read_only=True, data_only=True)
print("folhas:", livro.sheetnames)
folha = livro[livro.sheetnames[0]]
print("dimensoes:", folha.max_row, "x", folha.max_column)

linhas = []
for linha in folha.iter_rows(min_row=1, max_row=40, values_only=True):
    linhas.append(linha)
for numero, linha in enumerate(linhas, start=1):
    print(numero, "|", " | ".join("" if c is None else str(c)[:38] for c in linha))
